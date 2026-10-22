from __future__ import annotations

import base64
import io
import json
import zipfile
from email.message import EmailMessage
from pathlib import Path

import pytest
from PIL import Image

from ppt_king.documents import ConversionError, convert_document, supported_suffixes
from ppt_king.sources import scan_sources


def _png_bytes(size: tuple[int, int] = (80, 80)) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", size, (30, 90, 160)).save(buffer, format="PNG")
    return buffer.getvalue()


_DOCX_CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Default Extension="png" ContentType="image/png"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>
"""

_DOCX_ROOT_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>
"""

_DOCX_DOC_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image1.png"/>
</Relationships>
"""

_DOCX_DOCUMENT = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document
  xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
  xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
  xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
  xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>测试标题</w:t></w:r></w:p>
    <w:p><w:r><w:t>这是正文段落。</w:t></w:r></w:p>
    <w:p>
      <w:r><w:t>插图：</w:t></w:r>
      <w:r><w:drawing><a:graphic><a:graphicData><pic:pic><pic:blipFill>
        <a:blip r:embed="rId1"/>
      </pic:blipFill></pic:pic></a:graphicData></a:graphic></w:drawing></w:r>
    </w:p>
    <w:tbl>
      <w:tr><w:tc><w:p><w:r><w:t>指标</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>数值</w:t></w:r></w:p></w:tc></w:tr>
      <w:tr><w:tc><w:p><w:r><w:t>效率</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>81.2%</w:t></w:r></w:p></w:tc></w:tr>
    </w:tbl>
  </w:body>
</w:document>
"""


def _make_docx(path: Path) -> Path:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", _DOCX_CONTENT_TYPES)
        zf.writestr("_rels/.rels", _DOCX_ROOT_RELS)
        zf.writestr("word/document.xml", _DOCX_DOCUMENT)
        zf.writestr("word/_rels/document.xml.rels", _DOCX_DOC_RELS)
        zf.writestr("word/media/image1.png", _png_bytes())
    return path


def test_plain_and_code(tmp_path: Path) -> None:
    md = tmp_path / "a.md"
    md.write_text("# 标题\n\n正文\n", encoding="utf-8")
    result = convert_document(md)
    assert result.markdown.startswith("# 标题")

    code = tmp_path / "a.py"
    code.write_text("print('hi')\n", encoding="utf-8")
    result = convert_document(code)
    assert "```python" in result.markdown


def test_markdown_data_uri_image_is_extracted(tmp_path: Path) -> None:
    data = base64.b64encode(_png_bytes()).decode()
    md = tmp_path / "a.md"
    md.write_text(f"# 图\n\n![架构](data:image/png;base64,{data})\n", encoding="utf-8")

    result = convert_document(md)

    assert len(result.images) == 1
    assert "pptking-image://" in result.markdown


def test_csv_to_table(tmp_path: Path) -> None:
    path = tmp_path / "data.csv"
    path.write_text("指标,数值\n效率,81.2%\n转速,1200\n", encoding="utf-8")

    result = convert_document(path)

    assert "| 指标 | 数值 |" in result.markdown
    assert "| 效率 | 81.2% |" in result.markdown


def test_html_to_markdown_with_table_and_image(tmp_path: Path) -> None:
    data = base64.b64encode(_png_bytes()).decode()
    path = tmp_path / "page.html"
    path.write_text(
        "<html><body>"
        "<h1>架构</h1>"
        "<p>正文 <strong>加粗</strong></p>"
        "<ul><li>一</li><li>二</li></ul>"
        "<table><tr><th>项</th><th>值</th></tr><tr><td>A</td><td>1</td></tr></table>"
        f'<img src="data:image/png;base64,{data}" alt="示意图">'
        "</body></html>",
        encoding="utf-8",
    )

    result = convert_document(path)

    assert "# 架构" in result.markdown
    assert "**加粗**" in result.markdown
    assert "- 一" in result.markdown
    assert "| 项 | 值 |" in result.markdown
    assert len(result.images) == 1
    assert "![示意图](pptking-image://" in result.markdown


def test_docx_text_table_and_image(tmp_path: Path) -> None:
    path = _make_docx(tmp_path / "报告.docx")

    result = convert_document(path)

    assert "测试标题" in result.markdown
    assert "这是正文段落" in result.markdown
    assert "| 指标 | 数值 |" in result.markdown
    assert len(result.images) == 1
    assert "pptking-image://" in result.markdown


def test_pptx_text_table_and_image(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "source.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "测试页"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "效率"
    table.cell(1, 1).text = "81.2%"
    slide.shapes.add_picture(
        io.BytesIO(_png_bytes()), Inches(1), Inches(3), width=Inches(2), height=Inches(2)
    )
    prs.save(path)

    result = convert_document(path)

    assert "## 第 1 页：测试页" in result.markdown
    assert "| 指标 | 数值 |" in result.markdown
    assert len(result.images) == 1


def test_xlsx_to_tables(tmp_path: Path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "data.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "实验数据"
    ws.append(["指标", "数值"])
    ws.append(["效率", 0.812])
    ws.append(["数量", 3])
    wb.save(path)

    result = convert_document(path)

    assert "## 工作表：实验数据" in result.markdown
    assert "| 指标 | 数值 |" in result.markdown
    assert "| 效率 | 0.812 |" in result.markdown
    assert "| 数量 | 3 |" in result.markdown


def test_ipynb_cells_and_outputs(tmp_path: Path) -> None:
    data = base64.b64encode(_png_bytes((40, 40))).decode()
    notebook = {
        "metadata": {"kernelspec": {"language": "python"}},
        "cells": [
            {"cell_type": "markdown", "source": ["# 实验\n", "说明"]},
            {
                "cell_type": "code",
                "source": ["print('hi')"],
                "outputs": [
                    {"output_type": "stream", "text": ["hi\n"]},
                    {
                        "output_type": "display_data",
                        "data": {"image/png": data, "text/plain": ["<Figure>"]},
                    },
                ],
            },
        ],
    }
    path = tmp_path / "run.ipynb"
    path.write_text(json.dumps(notebook), encoding="utf-8")

    result = convert_document(path)

    assert "# 实验" in result.markdown
    assert "```python" in result.markdown
    assert "hi" in result.markdown
    assert len(result.images) == 1


def test_zip_converts_nested_sources(tmp_path: Path) -> None:
    path = tmp_path / "资料.zip"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("说明.md", "# 压缩包资料\n\n正文\n")
        zf.writestr("图.png", _png_bytes())

    result = convert_document(path)

    assert "## 文件：说明.md" in result.markdown
    assert "# 压缩包资料" in result.markdown
    assert len(result.images) == 1


def test_rtf_to_text(tmp_path: Path) -> None:
    path = tmp_path / "a.rtf"
    path.write_text(r"{\rtf1\ansi First paragraph\par Second paragraph}", encoding="latin-1")

    result = convert_document(path)

    assert "First paragraph" in result.markdown
    assert "Second paragraph" in result.markdown


def test_eml_subject_body_and_image(tmp_path: Path) -> None:
    message = EmailMessage()
    message["Subject"] = "汇报材料"
    message["From"] = "a@example.com"
    message["To"] = "b@example.com"
    message.set_content("正文内容")
    message.add_attachment(
        _png_bytes(), maintype="image", subtype="png", filename="figure.png"
    )
    path = tmp_path / "mail.eml"
    path.write_bytes(bytes(message))

    result = convert_document(path)

    assert "# 汇报材料" in result.markdown
    assert "正文内容" in result.markdown
    assert len(result.images) == 1


def test_unknown_suffix_raises(tmp_path: Path) -> None:
    path = tmp_path / "a.unknown"
    path.write_text("x", encoding="utf-8")
    with pytest.raises(ConversionError):
        convert_document(path)


def test_supported_suffixes_contains_documents() -> None:
    suffixes = set(supported_suffixes())
    assert {".pdf", ".docx", ".pptx", ".xlsx", ".html", ".zip", ".epub"} <= suffixes


def test_scan_sources_extracts_docx_text_and_images(tmp_path: Path) -> None:
    source = tmp_path / "资料"
    source.mkdir()
    (source / "outline.md").write_text("# 汇报\n\n## 封面\n", encoding="utf-8")
    _make_docx(source / "报告.docx")

    bundle = scan_sources(source, out_dir=tmp_path / "out")

    assert [doc.rel for doc in bundle.docs] == ["报告.docx"]
    doc = bundle.docs[0]
    assert "测试标题" in doc.text
    assert "assets/src01_" in doc.text
    assert bundle.images
    assert bundle.images[0].path.is_file()
    assert doc.title == "测试标题"
