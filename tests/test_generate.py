from __future__ import annotations

from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from ppt_king.config import Settings
from ppt_king.generate import GenerateError, default_output_dir, generate_deck
from ppt_king.grid import load_grid
from ppt_king.llm import LLMClient, LLMError


def _settings(tmp_path: Path) -> Settings:
    return Settings(
        api_key="k",
        base_url="http://localhost",
        model="fake",
        concurrency=1,
        cache_dir=tmp_path / "cache",
    )


def test_default_output_dir_is_out_name(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    source = tmp_path / "资料"
    source.mkdir()
    assert default_output_dir(source) == (tmp_path / "out" / "资料").resolve()


def _source(tmp_path: Path, *, with_outline: bool = True) -> Path:
    source = tmp_path / "资料"
    source.mkdir()
    if with_outline:
        (source / "outline.md").write_text(
            "# 我的汇报\n\n## 封面\n\n- 副标题\n\n## 内容\n\n- 要点\n",
            encoding="utf-8",
        )
    (source / "notes.md").write_text(
        "# 说明\n\n输入电压为 24V，输出 5V。\n\n![内容架构图](pic.png)\n",
        encoding="utf-8",
    )
    Image.new("RGB", (64, 36), (20, 60, 120)).save(source / "pic.png")
    return source


def _fake_complete(calls: list[str]):
    def complete(self, system: str, user: str) -> str:
        calls.append(system)
        if "资料分析员" in system:
            return "- 输入 24V，输出 5V"
        if "PPT 版式推荐器" in system:
            return '{"templates":[]}'
        if "大纲策划" in system:
            return (
                "# 自动标题\n\n## 封面\n\n- 副标题\n\n"
                "## 内容\n\n- 要点\n\n## 总结\n\n- 结论"
            )
        if "PPT 内容撰写器" in system:
            if "可用图片：无" in user:
                return "@ 2,3-15,8\n输入电压为 24V。"
            return (
                "@ 2,3-8,7\n输入电压为 24V。\n\n"
                "@ 10,3-15,7\n![内容架构图](assets/src01_pic.png)"
            )
        raise AssertionError(f"未知提示词：{system[:40]}")

    return complete


def test_generate_end_to_end(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    calls: list[str] = []
    monkeypatch.setattr(LLMClient, "complete", _fake_complete(calls))
    source = _source(tmp_path)
    out = tmp_path / "out"

    result = generate_deck(
        source,
        out=out,
        settings=_settings(tmp_path),
        min_pages=2,
        max_pages=3,
        progress=False,
        use_cache=False,
    )

    assert result.pptx_path.is_file()
    assert result.outline_path.is_file()
    assert result.grid_path.is_file()
    assert (out / "generate.json").is_file()
    assert (out / "pages" / "p01.md").is_file()
    assert (out / "pages" / "p02.md").is_file()
    assert (out / "assets" / "src01_pic.png").is_file()
    assert not result.failures
    assert len(calls) == 4  # 两页：每页推荐模板 + 写作

    deck = load_grid(result.grid_path)
    assert len(deck.slides) == 2
    prs = Presentation(str(result.pptx_path))
    assert len(prs.slides) == 2
    assert any(shape.shape_type == 13 for slide in prs.slides for shape in slide.shapes)


def test_generate_resume_skips_model(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    calls: list[str] = []
    monkeypatch.setattr(LLMClient, "complete", _fake_complete(calls))
    source = _source(tmp_path)
    out = tmp_path / "out"
    settings = _settings(tmp_path)

    generate_deck(source, out=out, settings=settings, progress=False, use_cache=False)
    calls.clear()
    result = generate_deck(
        source, out=out, settings=settings, resume=True, progress=False, use_cache=False
    )

    assert calls == []
    assert result.pptx_path.is_file()


def test_generate_force_rewrites_pages(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    calls: list[str] = []
    monkeypatch.setattr(LLMClient, "complete", _fake_complete(calls))
    source = _source(tmp_path)
    out = tmp_path / "out"
    settings = _settings(tmp_path)

    generate_deck(source, out=out, settings=settings, progress=False, use_cache=False)
    calls.clear()
    generate_deck(
        source,
        out=out,
        settings=settings,
        resume=True,
        force=True,
        progress=False,
        use_cache=False,
    )

    assert len(calls) == 4


def test_generate_ai_outline(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    calls: list[str] = []
    monkeypatch.setattr(LLMClient, "complete", _fake_complete(calls))
    source = _source(tmp_path, with_outline=False)
    out = tmp_path / "out"

    result = generate_deck(
        source,
        out=out,
        settings=_settings(tmp_path),
        min_pages=3,
        max_pages=3,
        progress=False,
        use_cache=False,
    )

    assert (out / "outline.md").is_file()
    assert result.pages == 3
    assert len(calls) == 8  # 摘要 + 大纲 + 三页 ×（推荐 + 写作）


def test_generate_adds_cover_when_ai_outline_omits_it(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    def fake_complete(self, system: str, user: str) -> str:
        if "资料分析员" in system:
            return "- 摘要"
        if "大纲策划" in system:
            return "# 自动标题\n\n## 背景\n\n- 要点\n\n## 总结\n\n- 结论"
        return "@ 2,3-15,8\n正文内容"

    monkeypatch.setattr(LLMClient, "complete", fake_complete)
    source = _source(tmp_path, with_outline=False)
    out = tmp_path / "out"

    result = generate_deck(
        source,
        out=out,
        settings=_settings(tmp_path),
        min_pages=2,
        max_pages=2,
        progress=False,
        use_cache=False,
    )

    assert result.pages == 3
    text = (out / "outline.md").read_text(encoding="utf-8")
    assert text.startswith("# 自动标题")
    assert "## 封面" in text


def test_generate_continues_after_page_failure(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    def fake_complete(self, system: str, user: str) -> str:
        if "页标题：内容" in user:
            raise LLMError("模拟单页失败")
        return "@ 2,3-15,8\n封面内容"

    monkeypatch.setattr(LLMClient, "complete", fake_complete)
    source = _source(tmp_path)
    out = tmp_path / "out"

    result = generate_deck(
        source,
        out=out,
        settings=_settings(tmp_path),
        progress=False,
        use_cache=False,
    )

    assert result.pptx_path.is_file()
    assert len(result.failures) == 1
    assert "模拟单页失败" in result.failures[0]
    assert "要点" in result.grid_path.read_text(encoding="utf-8")


def test_generate_strict_raises_on_failure(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    monkeypatch.setattr(
        LLMClient, "complete", lambda self, system, user: "这一页没有坐标"
    )
    source = _source(tmp_path)
    with pytest.raises(GenerateError):
        generate_deck(
            source,
            out=tmp_path / "out",
            settings=_settings(tmp_path),
            strict=True,
            progress=False,
            use_cache=False,
        )
