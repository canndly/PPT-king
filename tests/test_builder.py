from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu

from ppt_king.builder import SLIDE_H, SLIDE_W, build_deck
from ppt_king.rebuild import make_theme as rebuild_make_theme
from ppt_king.spec import Block, BulletItem, DeckSpec, SlideSpec


def _png(path: Path, size=(320, 200)) -> Path:
    from PIL import Image

    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, (30, 90, 200)).save(path)
    return path


@pytest.fixture()
def deck_spec(tmp_path: Path) -> DeckSpec:
    assets = tmp_path / "assets"
    _png(assets / "a.png", (600, 300))
    _png(assets / "b.png", (300, 600))
    return DeckSpec(
        title="还原测试",
        slides=[
            SlideSpec(
                number=1,
                layout="title",
                title="课程设计报告",
                blocks=[
                    Block(type="paragraph", text="基于 LM5180 的反激电路设计"),
                    Block(type="image", ref="A1", caption="封面图"),
                ],
                assets={"A1": "assets/a.png"},
            ),
            SlideSpec(
                number=2,
                layout="two_column",
                title="一、题目解析",
                blocks=[
                    Block(type="heading", text="设计要求"),
                    Block(type="paragraph", text="输入 24 V，输出 5 V/1 A。"),
                    Block(
                        type="bullets",
                        items=[BulletItem("隔离型电源", 0), BulletItem("效率优先", 1)],
                    ),
                    Block(type="image", ref="A1", caption="原理图"),
                ],
                assets={"A1": "assets/a.png"},
                notes="这是备注",
            ),
            SlideSpec(
                number=3,
                layout="table",
                title="实验数据",
                blocks=[
                    Block(
                        type="table",
                        header=["输入电压", "输出电压", "效率"],
                        rows=[["24 V", "5.036 V", "81.24%"], ["24 V", "5.081 V", "82.01%"]],
                    )
                ],
            ),
            SlideSpec(
                number=4,
                layout="gallery",
                title="波形展示",
                blocks=[
                    Block(type="image", ref="A1", caption="上电波形"),
                    Block(type="image", ref="A2", caption="纹波波形"),
                ],
                assets={"A1": "assets/a.png", "A2": "assets/b.png"},
            ),
            SlideSpec(number=5, layout="section", title="七、心得体会"),
        ],
    )


def test_build_deck_structure(deck_spec: DeckSpec, tmp_path: Path) -> None:
    out = build_deck(deck_spec, tmp_path / "out.pptx", base_dir=tmp_path)
    assert out.is_file()
    prs = Presentation(str(out))
    assert len(prs.slides) == 5
    assert prs.slide_width == Emu(SLIDE_W)
    assert prs.slide_height == Emu(SLIDE_H)
    assert prs.core_properties.title == "还原测试"


def test_pictures_placed(deck_spec: DeckSpec, tmp_path: Path) -> None:
    out = build_deck(deck_spec, tmp_path / "out.pptx", base_dir=tmp_path)
    prs = Presentation(str(out))

    def picture_count(slide) -> int:
        return sum(1 for shape in slide.shapes if shape.shape_type == 13)

    assert picture_count(prs.slides[0]) >= 1
    assert picture_count(prs.slides[1]) >= 1
    assert picture_count(prs.slides[3]) == 2


def test_table_content(deck_spec: DeckSpec, tmp_path: Path) -> None:
    out = build_deck(deck_spec, tmp_path / "out.pptx", base_dir=tmp_path)
    prs = Presentation(str(out))
    tables = [shape for shape in prs.slides[2].shapes if shape.has_table]
    assert len(tables) == 1
    table = tables[0].table
    assert len(table.rows) == 3
    assert table.cell(0, 0).text == "输入电压"
    assert table.cell(2, 2).text == "82.01%"


def test_text_and_notes(deck_spec: DeckSpec, tmp_path: Path) -> None:
    out = build_deck(deck_spec, tmp_path / "out.pptx", base_dir=tmp_path)
    prs = Presentation(str(out))
    slide2_text = " ".join(
        shape.text_frame.text for shape in prs.slides[1].shapes if shape.has_text_frame
    )
    assert "题目解析" in slide2_text
    assert "隔离型电源" in slide2_text
    assert prs.slides[1].notes_slide.notes_text_frame.text == "这是备注"


def test_slide_numbers_rendered(deck_spec: DeckSpec, tmp_path: Path) -> None:
    out = build_deck(deck_spec, tmp_path / "out.pptx", base_dir=tmp_path)
    prs = Presentation(str(out))
    texts = [shape.text_frame.text for shape in prs.slides[2].shapes if shape.has_text_frame]
    assert "3" in texts


def test_cover_has_no_page_number(deck_spec: DeckSpec, tmp_path: Path) -> None:
    out = build_deck(deck_spec, tmp_path / "out.pptx", base_dir=tmp_path)
    prs = Presentation(str(out))
    texts = [shape.text_frame.text for shape in prs.slides[0].shapes if shape.has_text_frame]
    assert "1" not in texts


def test_missing_image_is_skipped(deck_spec: DeckSpec, tmp_path: Path) -> None:
    deck_spec.slides[3].assets = {}
    out = build_deck(deck_spec, tmp_path / "out.pptx", base_dir=tmp_path)
    prs = Presentation(str(out))
    assert len(prs.slides) == 5  # 不因缺图崩溃


def test_shapes_stay_inside_slide(deck_spec: DeckSpec, tmp_path: Path) -> None:
    out = build_deck(deck_spec, tmp_path / "out.pptx", base_dir=tmp_path)
    prs = Presentation(str(out))
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left is not None and shape.top is not None
            assert shape.left >= -Emu(10000)
            assert shape.top >= -Emu(10000)
            assert shape.left + shape.width <= Emu(SLIDE_W + 10000)
            assert shape.top + shape.height <= Emu(SLIDE_H + 10000)


def test_accent_theme_override() -> None:
    theme = rebuild_make_theme("FF0000")
    assert str(theme.accent) == "FF0000"


def test_invalid_accent_rejected() -> None:
    from ppt_king.spec import SpecError

    with pytest.raises(SpecError):
        rebuild_make_theme("red")
