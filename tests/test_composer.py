from __future__ import annotations

import json
from pathlib import Path

import pytest

from ppt_king.composer import PageSource, build_user_message, split_markdown
from ppt_king.llm import LLMError
from ppt_king.rebuild import default_output, locate_markdown, rebuild_from_export, rebuild_from_spec

MARKDOWN = """# 电力电子课程设计报告

## 1. 电力电子课程设计报告

**基于 LM5180 的反激电路设计**

指导老师：郑峰
姓名：李泽海

![封面](assets/cover.png)

---

## 2. 一、题目解析

本设计要求基于 LM5180 设计隔离电源。

![原理图](assets/circuit.png)

> **备注**：这里是备注第一行
> 备注第二行

---

## 3. 六，波形展示

### 图表：效率曲线

| 分类 | 效率 |
| --- | --- |
| 0.2A | 80.1 |
"""


@pytest.fixture()
def export_dir(tmp_path: Path) -> Path:
    from PIL import Image

    root = tmp_path / "export" / "示例"
    (root / "assets").mkdir(parents=True)
    Image.new("RGB", (400, 200), (10, 20, 30)).save(root / "assets" / "cover.png")
    Image.new("RGB", (200, 400), (40, 50, 60)).save(root / "assets" / "circuit.png")
    (root / "示例.md").write_text(MARKDOWN, encoding="utf-8")
    return root


def test_locate_markdown_from_dir(export_dir: Path) -> None:
    assert locate_markdown(export_dir) == export_dir / "示例.md"


def test_default_output(export_dir: Path) -> None:
    md = export_dir / "示例.md"
    assert default_output(md) == export_dir / "示例.pptx"


def test_split_markdown_pages(export_dir: Path) -> None:
    split = split_markdown(MARKDOWN, export_dir)
    assert split.title == "电力电子课程设计报告"
    assert [page.number for page in split.pages] == [1, 2, 3]
    assert split.pages[1].title == "一、题目解析"
    assert split.pages[2].title == "六，波形展示"


def test_split_markdown_indexes_images(export_dir: Path) -> None:
    split = split_markdown(MARKDOWN, export_dir)
    cover = split.pages[0]
    assert list(cover.assets) == ["A1"]
    assert cover.assets["A1"] == "assets/cover.png"
    assert "{{A1}}" in cover.body
    assert "assets/cover.png" not in cover.body
    assert cover.captions["A1"] == "封面"


def test_split_markdown_extracts_notes(export_dir: Path) -> None:
    page = split_markdown(MARKDOWN, export_dir).pages[1]
    assert "备注第一行" in page.notes
    assert "备注第二行" in page.notes
    assert "备注" not in page.body


def test_split_markdown_keeps_unknown_images(export_dir: Path) -> None:
    text = "## 1. 页\n\n![缺失](assets/none.png)"
    page = split_markdown(text, export_dir).pages[0]
    assert page.assets == {}


def test_build_user_message_lists_refs(export_dir: Path) -> None:
    page = split_markdown(MARKDOWN, export_dir).pages[1]
    message = build_user_message(page, 3, export_dir)
    assert "A1 = assets/circuit.png" in message
    assert "图注=原理图" in message
    assert "像素=200x400" in message


def test_split_markdown_requires_headings(tmp_path: Path) -> None:
    from ppt_king.spec import SpecError

    with pytest.raises(SpecError):
        split_markdown("# 只有标题\n正文", tmp_path)


def test_compose_and_build(export_dir: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from ppt_king.composer import compose_deck
    from ppt_king.config import Settings
    from ppt_king.llm import LLMClient

    responses = {
        "幻灯片 1/3": {
            "layout": "title",
            "blocks": [
                {"type": "paragraph", "text": "基于 LM5180 的反激电路设计"},
                {"type": "paragraph", "text": "指导老师：郑峰"},
                {"type": "image", "ref": "A1", "caption": "封面"},
            ],
        },
        "幻灯片 2/3": {
            "layout": "two_column",
            "blocks": [
                {"type": "paragraph", "text": "本设计要求基于 LM5180 设计隔离电源。"},
                {"type": "image", "ref": "A1", "caption": "原理图"},
            ],
        },
        "幻灯片 3/3": {
            "layout": "table",
            "blocks": [
                {"type": "paragraph", "text": "六，波形展示"},
                {"type": "table", "header": ["分类", "效率"], "rows": [["0.2A", "80.1"]]},
            ],
        },
    }

    def fake_complete(self, system: str, user: str) -> str:
        for prefix, payload in responses.items():
            if user.startswith(prefix):
                return "```json\n" + json.dumps(payload) + "\n```"
        raise LLMError("未覆盖")

    monkeypatch.setattr(LLMClient, "complete", fake_complete)

    settings = Settings(api_key="k", base_url="http://localhost", model="fake", concurrency=2)
    client = LLMClient(settings)
    from ppt_king.cache import Cache

    cache = Cache.create(export_dir / "cache", enabled=False)
    spec = compose_deck(
        export_dir / "示例.md",
        settings=settings,
        client=client,
        cache=cache,
        concurrency=2,
        progress=False,
    )
    assert len(spec.slides) == 3
    assert spec.slides[0].layout == "title"
    assert spec.slides[1].assets == {"A1": "assets/circuit.png"}
    assert spec.slides[2].blocks[-1].header == ["分类", "效率"]


def test_invalid_json_triggers_retry(export_dir: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from ppt_king.cache import Cache
    from ppt_king.composer import compose_slide
    from ppt_king.config import Settings
    from ppt_king.llm import LLMClient

    calls: list[str] = []

    def fake_complete(self, system: str, user: str) -> str:
        calls.append(user)
        if len(calls) == 1:
            return "不是 JSON"
        return '{"layout": "section", "blocks": []}'

    monkeypatch.setattr(LLMClient, "complete", fake_complete)

    settings = Settings(api_key="k", base_url="http://localhost", model="fake")
    client = LLMClient(settings)
    page = PageSource(number=1, title="章节", body="内容")
    slide = compose_slide(
        client, Cache.create(export_dir / "c", enabled=False), settings, page, 1, export_dir
    )
    assert len(calls) == 2
    assert slide.layout == "section"
    assert slide.title == "章节"


def test_unknown_ref_triggers_retry(export_dir: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from ppt_king.cache import Cache
    from ppt_king.composer import compose_slide
    from ppt_king.config import Settings
    from ppt_king.llm import LLMClient

    outputs = [
        '{"layout": "title_content", "blocks": [{"type": "image", "ref": "A7"}]}',
        '{"layout": "title_content", "blocks": [{"type": "image", "ref": "A1"}]}',
    ]

    def fake_complete(self, system: str, user: str) -> str:
        return outputs.pop(0)

    monkeypatch.setattr(LLMClient, "complete", fake_complete)

    settings = Settings(api_key="k", base_url="http://localhost", model="fake")
    page = PageSource(
        number=1,
        title="页",
        body="{{A1}}",
        assets={"A1": "assets/a.png"},
        captions={"A1": ""},
    )
    slide = compose_slide(
        LLMClient(settings),
        Cache.create(export_dir / "c", enabled=False),
        settings,
        page,
        1,
        export_dir,
    )
    assert slide.blocks[0].ref == "A1"


def test_rebuild_end_to_end(export_dir: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from ppt_king.config import Settings
    from ppt_king.llm import LLMClient

    payloads = {
        "幻灯片 1/3": '{"layout": "title", "blocks": [{"type": "image", "ref": "A1"}]}',
        "幻灯片 2/3": '{"layout": "two_column", "blocks": [{"type": "paragraph", "text": "正文"}, {"type": "image", "ref": "A1"}]}',
        "幻灯片 3/3": '{"layout": "table", "blocks": [{"type": "table", "header": ["a", "b"], "rows": [["1", "2"]]}]}',
    }

    def fake_complete(self, system: str, user: str) -> str:
        for prefix, payload in payloads.items():
            if user.startswith(prefix):
                return payload
        raise LLMError("未覆盖")

    monkeypatch.setattr(LLMClient, "complete", fake_complete)

    settings = Settings(api_key="k", base_url="http://localhost", model="fake", concurrency=1)
    result = rebuild_from_export(
        export_dir,
        settings=settings,
        use_cache=False,
        save_spec=export_dir / "spec.json",
        progress=False,
    )
    assert result.output.is_file()
    assert result.output.name == "示例.pptx"
    assert result.spec_path is not None and result.spec_path.is_file()

    from pptx import Presentation

    prs = Presentation(str(result.output))
    assert len(prs.slides) == 3


def test_rebuild_from_saved_spec(export_dir: Path) -> None:
    from PIL import Image

    from ppt_king.spec import Block, DeckSpec, SlideSpec

    spec = DeckSpec(
        title="t",
        slides=[
            SlideSpec(
                number=1,
                layout="title_content",
                title="页",
                blocks=[Block(type="image", ref="A1", caption="图")],
                assets={"A1": "assets/cover.png"},
            )
        ],
    )
    spec_path = spec.save(export_dir / "spec.json")
    result = rebuild_from_spec(spec_path, out=export_dir / "from-spec.pptx", progress=False)
    assert result.output.is_file()
    assert Image.open(export_dir / "assets" / "cover.png").size == (400, 200)
