from __future__ import annotations

from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

from ppt_king.grid import (
    CODE,
    EMU_PER_CELL,
    FORMULA,
    GRID_COLS,
    GRID_ROWS,
    IMAGE,
    QUOTE,
    TABLE,
    TEXT,
    Box,
    GridSyntaxError,
    describe,
    load_grid,
    parse_grid,
)
from ppt_king.grid_render import SLIDE_H, SLIDE_W, _segments, make_theme, render_deck

SOURCE = """# 演示文稿标题

## 第一页

@ 1,2 8x4
第一段文字。

第二段文字。

## 第二页

@ 1,2 8x3
引导语：

1. 第一条
2. 第二条

@ 10,2 7x5
![原理图](assets/a.png)

@ 1,6 16x3
| 名称 | 数值 |
| --- | --- |
| 电压 | 5 V |
"""


# --------------------------------------------------------------------------
# 解析
# --------------------------------------------------------------------------
def test_deck_title_and_pages() -> None:
    deck = parse_grid(SOURCE)
    assert deck.title == "演示文稿标题"
    assert [slide.title for slide in deck.slides] == ["第一页", "第二页"]
    assert [slide.number for slide in deck.slides] == [1, 2]


def test_blank_line_does_not_split_block() -> None:
    deck = parse_grid(SOURCE)
    block = deck.slides[0].blocks[0]
    assert block.kind == TEXT
    assert [item.text for item in block.flow] == ["第一段文字。", "第二段文字。"]


def test_paragraph_and_list_live_in_one_block() -> None:
    deck = parse_grid(SOURCE)
    block = deck.slides[1].blocks[0]
    kinds = [item.kind for item in block.flow]
    assert kinds == ["para", "ordered", "ordered"]
    assert block.flow[0].text == "引导语："


def test_anchor_with_span_align_and_size() -> None:
    deck = parse_grid("## 页\n\n@ 3,4 7x2 center 24\n正文\n")
    box = deck.slides[0].blocks[0].box
    assert box is not None
    assert (box.col, box.row, box.colspan, box.rowspan) == (3, 4, 7, 2)
    assert box.align == "center"
    assert box.size == 24


def test_anchor_range_coordinates() -> None:
    deck = parse_grid("## 页\n\n@ 3,4-9,5 center 24\n正文\n")
    box = deck.slides[0].blocks[0].box
    assert box is not None
    assert (box.col, box.row, box.colspan, box.rowspan) == (3, 4, 7, 2)
    assert box.align == "center"
    assert box.size == 24
    assert box.width == 7 * EMU_PER_CELL
    assert box.height == 2 * EMU_PER_CELL


def test_anchor_range_rejects_inverted_corner() -> None:
    with pytest.raises(GridSyntaxError, match="坐标范围非法"):
        parse_grid("## 页\n\n@ 5,5-4,6\n正文\n")


def test_anchor_without_span_extends_to_edge() -> None:
    deck = parse_grid("## 页\n\n@ 5,2\n正文\n")
    box = deck.slides[0].blocks[0].box
    assert box is not None
    assert box.colspan is None
    assert box.width == (GRID_COLS - 5 + 1) * EMU_PER_CELL
    assert box.height == (GRID_ROWS - 2 + 1) * EMU_PER_CELL


def test_out_of_range_is_clamped() -> None:
    deck = parse_grid("## 页\n\n@ 15,8 9x9\n正文\n")
    box = deck.slides[0].blocks[0].box
    assert box is not None
    assert box.colspan == GRID_COLS - 15 + 1
    assert box.rowspan == GRID_ROWS - 8 + 1
    assert box.x + box.width <= SLIDE_W
    assert box.y + box.height <= SLIDE_H


def test_table_block() -> None:
    deck = parse_grid(SOURCE)
    table = next(b for b in deck.slides[1].blocks if b.kind == TABLE)
    assert table.header == ["名称", "数值"]
    assert table.rows == [["电压", "5 V"]]


def test_image_block_uses_alt_as_caption() -> None:
    deck = parse_grid(SOURCE)
    image = next(b for b in deck.slides[1].blocks if b.kind == IMAGE)
    assert image.path == "assets/a.png"
    assert image.caption == "原理图"


def test_image_caption_prefix() -> None:
    deck = parse_grid("## 页\n\n@ 1,2 8x4\n![图](a.png)\n图注：仿真波形\n")
    image = deck.slides[0].blocks[0]
    assert image.caption == "仿真波形"


def test_quote_and_code_blocks() -> None:
    deck = parse_grid('## 页\n\n@ 1,2 8x2\n> 引用内容\n\n@ 1,5 8x3\n```\nx = 1\n```\n')
    kinds = [block.kind for block in deck.slides[0].blocks]
    assert kinds == [QUOTE, CODE]
    assert deck.slides[0].blocks[1].flow[0].text == "x = 1"


def test_cover_subtitle_with_dot_is_not_stripped() -> None:
    deck = parse_grid(
        "## 封面\n\n@ 1,4 16x1 center 34\n主标题\n\n"
        "@ 1,6 16x1 center 18\n课程答辩 · 零样本迁移\n"
    )
    assert len(deck.slides[0].blocks) == 2
    assert deck.slides[0].blocks[1].flow[0].text == "课程答辩 · 零样本迁移"


def test_formula_block_parse() -> None:
    deck = parse_grid(
        "## 页\n\n@ 1,3 16x1 center 24\n$$ v·ω ≤ μg $$\n"
    )
    block = deck.slides[0].blocks[0]
    assert block.kind == FORMULA
    assert block.flow[0].text == "v·ω ≤ μg"


def test_headless_page() -> None:
    deck = parse_grid("##\n@ 1,3 16x2 center 34\nThanks\n")
    assert deck.slides[0].title == ""
    assert deck.slides[0].blocks[0].box is not None


def test_positioned_without_anchor_goes_auto() -> None:
    deck = parse_grid("## 页\n\n未定位\n\n@ 1,5 8x3\n定位\n")
    assert not deck.slides[0].blocks[0].is_positioned
    assert deck.slides[0].blocks[1].is_positioned


def test_content_without_heading_becomes_one_page() -> None:
    deck = parse_grid("只有正文，没有 `##` 页标题")
    assert len(deck.slides) == 1
    assert deck.slides[0].title == ""


def test_empty_input_raises() -> None:
    with pytest.raises(GridSyntaxError):
        parse_grid("\n\n   \n")


def test_describe_lists_blocks() -> None:
    text = describe(parse_grid(SOURCE))
    assert "第一页" in text
    assert "@1,2-8,5" in text
    assert "table" in text


# --------------------------------------------------------------------------
# 渲染
# --------------------------------------------------------------------------
def _render(
    md: str,
    tmp_path: Path,
    *,
    logo: Path | bool | None = False,
    header: str | None = None,
) -> tuple[Path, object]:
    (tmp_path / "assets").mkdir(exist_ok=True)
    Image.new("RGB", (640, 360), (20, 60, 120)).save(tmp_path / "assets" / "a.png")
    source = tmp_path / "deck.md"
    source.write_text(md, encoding="utf-8")
    deck = load_grid(source)
    return render_deck(
        deck,
        tmp_path / "deck.pptx",
        base_dir=tmp_path,
        logo=logo,
        header=header,
    )


def test_render_widens_lone_text_block(tmp_path: Path) -> None:
    out, _ = _render("## 页\n\n@ 1,3-8,8\n正文内容\n", tmp_path)
    prs = Presentation(str(out))
    body = next(
        shape
        for shape in prs.slides[0].shapes
        if shape.has_text_frame and "正文内容" in shape.text_frame.text
    )
    assert body.width == SLIDE_W - 2 * EMU_PER_CELL


def test_render_creates_slides(tmp_path: Path) -> None:
    out, report = _render(SOURCE, tmp_path)
    prs = Presentation(str(out))
    assert len(prs.slides) == 2
    assert prs.slide_width == Emu(SLIDE_W)
    assert prs.slide_height == Emu(SLIDE_H)
    assert report.slides == 2
    assert report.blocks == 4


def test_render_shapes_inside_slide(tmp_path: Path) -> None:
    out, _ = _render(SOURCE, tmp_path)
    prs = Presentation(str(out))
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= SLIDE_W + 1000
            assert shape.top + shape.height <= SLIDE_H + 1000


def test_render_table_cells(tmp_path: Path) -> None:
    out, _ = _render(SOURCE, tmp_path)
    prs = Presentation(str(out))
    table = next(s.table for s in prs.slides[1].shapes if s.has_table)
    assert table.cell(0, 0).text == "名称"
    assert table.cell(1, 1).text == "5 V"


def test_render_picture_placed(tmp_path: Path) -> None:
    out, _ = _render(SOURCE, tmp_path)
    prs = Presentation(str(out))
    pictures = [s for s in prs.slides[1].shapes if s.shape_type == 13]
    assert len(pictures) == 1


def test_render_text_and_markers(tmp_path: Path) -> None:
    out, _ = _render(SOURCE, tmp_path)
    prs = Presentation(str(out))
    texts = [
        shape.text_frame.text for shape in prs.slides[1].shapes if shape.has_text_frame
    ]
    joined = "\n".join(texts)
    assert "引导语：" in joined
    assert "1. 第一条" in joined
    assert "2. 第二条" in joined


def test_render_missing_image_reports(tmp_path: Path) -> None:
    source = tmp_path / "m.md"
    source.write_text("## 页\n\n@ 1,2 8x4\n![x](assets/none.png)\n", encoding="utf-8")
    deck = load_grid(source)
    _, report = render_deck(deck, tmp_path / "m.pptx", base_dir=tmp_path)
    assert any("assets/none.png" in item for item in report.missing_images)
    assert not report.overflows


def test_render_auto_block_when_no_anchor(tmp_path: Path) -> None:
    source = tmp_path / "auto.md"
    source.write_text("## 页\n\n没有坐标的正文\n", encoding="utf-8")
    deck = load_grid(source)
    _, report = render_deck(deck, tmp_path / "auto.pptx", base_dir=tmp_path)
    assert report.auto_blocks == 1


def test_render_cover_skips_top_title(tmp_path: Path) -> None:
    out, _ = _render(
        "## 封面\n\n@ 1,4 16x1 center 34\n主标题\n\n"
        "@ 1,6 16x1 center 18\n副标题\n",
        tmp_path,
    )
    prs = Presentation(str(out))
    texts = [
        shape.text_frame.text
        for shape in prs.slides[0].shapes
        if shape.has_text_frame
    ]
    assert "封面" not in texts
    assert "主标题" in texts
    assert "副标题" in texts


def test_inline_segments_strip_backticks_and_mark_red() -> None:
    parts = _segments("前 `公式` <red>重点</red> **粗体**")
    assert "".join(part.text for part in parts) == "前 公式 重点 粗体"
    assert [part.text for part in parts if part.red] == ["重点"]
    assert [part.text for part in parts if part.bold] == ["粗体"]


def test_render_inline_red_and_backticks(tmp_path: Path) -> None:
    out, _ = _render(
        "## 页\n\n@ 1,2 16x3\n前 `公式` 和 <red>重点</red>\n", tmp_path
    )
    prs = Presentation(str(out))
    runs = [
        run
        for shape in prs.slides[0].shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    ]
    text = "".join(run.text for run in runs)
    assert "`" not in text
    assert "公式" in text
    red = [
        run
        for run in runs
        if run.text == "重点" and str(run.font.color.rgb) == "FF0000"
    ]
    assert red


def test_theme_override() -> None:
    theme = make_theme("C00000")
    assert str(theme.table_head) == "C00000"
    with pytest.raises(ValueError):
        make_theme("红色")


def test_box_geometry() -> None:
    box = Box(2, 3, 4, 2)
    assert box.x == EMU_PER_CELL
    assert box.y == 2 * EMU_PER_CELL
    assert box.width == 4 * EMU_PER_CELL
    assert box.height == 2 * EMU_PER_CELL


def test_box_bottom_limit() -> None:
    box = Box(1, 2, 16, 8)
    limited = box.with_bottom_limit(SLIDE_H - EMU_PER_CELL)
    assert limited.y + limited.height == SLIDE_H - EMU_PER_CELL


def test_render_keeps_content_above_bottom_row(tmp_path: Path) -> None:
    out, _ = _render("## 页\n\n@ 1,2 16x8\n正文\n", tmp_path)
    prs = Presentation(str(out))
    for shape in prs.slides[0].shapes:
        assert shape.top + shape.height <= SLIDE_H - EMU_PER_CELL + 1000


def test_render_logo_every_slide_and_title_shift(tmp_path: Path) -> None:
    logo = tmp_path / "logo.png"
    Image.new("RGB", (200, 50), (180, 0, 0)).save(logo)
    out, _ = _render(
        "## 页一\n\n@ 1,2 16x3\n正文一\n\n## 页二\n\n@ 1,2 16x3\n正文二\n",
        tmp_path,
        logo=logo,
        header="页眉",
    )
    prs = Presentation(str(out))
    for slide in prs.slides:
        pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
        assert len(pictures) == 1
        picture = pictures[0]
        assert picture.left == 0 and picture.top >= 0
        assert picture.top + picture.height <= EMU_PER_CELL + 1000
        assert picture.left + picture.width <= 8 * EMU_PER_CELL + 1000

    first = list(prs.slides)[0]
    title = next(
        shape
        for shape in first.shapes
        if shape.has_text_frame and "页一" in shape.text_frame.text
    )
    assert title.top >= EMU_PER_CELL

    header = next(
        shape
        for shape in first.shapes
        if shape.has_text_frame and "页眉" in shape.text_frame.text
    )
    assert header.left >= 8 * EMU_PER_CELL
    assert header.top == 0
    assert header.height <= EMU_PER_CELL + 1000

    rule = next(
        shape
        for shape in first.shapes
        if shape.height == 9000 and shape.width == SLIDE_W
    )
    assert rule.top == EMU_PER_CELL - 9000


def test_all_text_and_tables_at_least_15pt(tmp_path: Path) -> None:
    out, _ = _render(
        "## 页\n\n@ 1,3 16x3\n正文内容\n\n"
        "@ 1,6 16x3\n| 列一 | 列二 |\n| --- | --- |\n| 数值 | 20 |\n",
        tmp_path,
    )
    prs = Presentation(str(out))
    sizes: list[float] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            runs = []
            if shape.has_text_frame:
                runs.extend(
                    run
                    for paragraph in shape.text_frame.paragraphs
                    for run in paragraph.runs
                )
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        runs.extend(
                            run
                            for paragraph in cell.text_frame.paragraphs
                            for run in paragraph.runs
                        )
            sizes.extend(
                run.font.size.pt
                for run in runs
                if run.text.strip() and run.font.size is not None
            )
    assert sizes
    assert min(sizes) >= 15
    assert max(set(sizes), key=sizes.count) == 20


def test_default_logo_is_page_logo() -> None:
    from ppt_king.grid_render import DEFAULT_LOGO

    assert DEFAULT_LOGO.name == "page_logo.png"
    assert DEFAULT_LOGO.is_file()


def test_render_formula_block_italic(tmp_path: Path) -> None:
    out, _ = _render(
        "## 页\n\n@ 1,3 16x1 center 24\n$$ v·ω ≤ μg $$\n", tmp_path
    )
    prs = Presentation(str(out))
    runs = [
        run
        for shape in prs.slides[0].shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if run.text.strip()
    ]
    formula = next(run for run in runs if "v·ω" in run.text)
    assert formula.font.italic
    assert formula.font.size.pt >= 20


def test_render_formula_normalizes_latex(tmp_path: Path) -> None:
    out, _ = _render(
        "## 页\n\n@ 1,3 16x1 center 24\n"
        "$$ m·v·\\omega \\le \\mu mg \\quad\\Rightarrow\\quad "
        "v·\\omega \\le \\mu g $$\n",
        tmp_path,
    )
    prs = Presentation(str(out))
    runs = [
        run
        for shape in prs.slides[0].shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if run.text.strip()
    ]
    formula = next(run for run in runs if "m·v" in run.text)
    assert "\\" not in formula.text
    assert "\\quad" not in formula.text
    assert "⇒" in formula.text
    assert "≤" in formula.text
    assert "ω" in formula.text
    assert "μ" in formula.text
