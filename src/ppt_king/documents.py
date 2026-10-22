"""多格式资料转换：把常见文档转成 Markdown，并抽出内嵌图片。

设计参考 markitdown 的 converter 思路：每种后缀对应一个转换器，输出统一为
Markdown；图片不写成 base64，而是抽出为 :class:`EmbeddedImage`，由上游落盘到
``assets/``，这样后续逐页写作可以把它们当作可用配图。

支持范围（缺失依赖时会给出可读的 :class:`ConversionError`）：

* 纯文本 / Markdown / 代码 / JSON / XML / YAML / RTF
* CSV / TSV
* HTML
* Word ``.docx``
* PowerPoint ``.pptx``
* Excel ``.xlsx`` / ``.xls``
* PDF
* EPUB
* Jupyter Notebook ``.ipynb``
* 邮件 ``.eml``
* ZIP（递归转换其中的资料）
"""

from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import posixpath
import re
import tempfile
import zipfile
from dataclasses import dataclass, field
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from pathlib import Path
from typing import Callable
from urllib.parse import unquote, urlparse
from xml.etree import ElementTree as ET

MAX_TEXT_BYTES = 2_000_000
MAX_TABLE_ROWS = 500
MAX_TABLE_COLS = 40
MAX_ZIP_MEMBERS = 200
MAX_PDF_PAGES = 300
MAX_DOC_CHARS = 400_000

IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".svg", ".tif", ".tiff"}
PLAIN_SUFFIXES = {
    ".md",
    ".markdown",
    ".txt",
    ".text",
    ".log",
    ".rst",
    ".org",
    ".tex",
    ".srt",
    ".vtt",
}
CODE_SUFFIXES = {
    ".py",
    ".js",
    ".mjs",
    ".ts",
    ".tsx",
    ".jsx",
    ".java",
    ".c",
    ".h",
    ".cpp",
    ".hpp",
    ".cc",
    ".cs",
    ".go",
    ".rs",
    ".rb",
    ".php",
    ".swift",
    ".kt",
    ".scala",
    ".sh",
    ".bash",
    ".zsh",
    ".ps1",
    ".sql",
    ".css",
    ".scss",
    ".vue",
    ".ini",
    ".cfg",
    ".conf",
    ".toml",
    ".yaml",
    ".yml",
    ".xml",
    ".json",
    ".jsonl",
    ".r",
    ".m",
    ".lua",
}
HTML_SUFFIXES = {".html", ".htm", ".xhtml"}
DOC_SUFFIXES = {
    ".docx",
    ".pptx",
    ".xlsx",
    ".xls",
    ".pdf",
    ".epub",
    ".ipynb",
    ".rtf",
    ".zip",
    ".eml",
}
SUPPORTED_SUFFIXES = frozenset(
    PLAIN_SUFFIXES | CODE_SUFFIXES | HTML_SUFFIXES | DOC_SUFFIXES | {".csv", ".tsv"}
)
PLAIN_TEXT_SUFFIXES = frozenset({".md", ".markdown", ".txt", ".text"})

_DATA_URI_RE = re.compile(r"^data:([^;,]*)(;base64)?,(.*)$", re.S | re.I)
_MD_DATA_IMAGE_RE = re.compile(r"!\[([^\]]*)\]\((data:[^)\s]+)\)", re.I)
_MIME_SUFFIX = {
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/jpg": ".jpg",
    "image/gif": ".gif",
    "image/bmp": ".bmp",
    "image/webp": ".webp",
    "image/svg+xml": ".svg",
    "image/tiff": ".tiff",
}
_LANG = {
    ".py": "python",
    ".js": "javascript",
    ".mjs": "javascript",
    ".ts": "typescript",
    ".tsx": "tsx",
    ".jsx": "jsx",
    ".java": "java",
    ".c": "c",
    ".h": "c",
    ".cpp": "cpp",
    ".hpp": "cpp",
    ".cc": "cpp",
    ".cs": "csharp",
    ".go": "go",
    ".rs": "rust",
    ".rb": "ruby",
    ".php": "php",
    ".swift": "swift",
    ".kt": "kotlin",
    ".scala": "scala",
    ".sh": "bash",
    ".bash": "bash",
    ".zsh": "bash",
    ".ps1": "powershell",
    ".sql": "sql",
    ".css": "css",
    ".scss": "scss",
    ".vue": "vue",
    ".ini": "ini",
    ".cfg": "ini",
    ".conf": "ini",
    ".toml": "toml",
    ".yaml": "yaml",
    ".yml": "yaml",
    ".xml": "xml",
    ".json": "json",
    ".jsonl": "json",
    ".r": "r",
    ".m": "matlab",
    ".lua": "lua",
}


class ConversionError(ValueError):
    """资料格式无法转换，或缺少可选依赖。"""


@dataclass(frozen=True)
class EmbeddedImage:
    key: str
    data: bytes
    suffix: str
    alt: str = ""
    title: str = ""


@dataclass
class ConvertedDocument:
    markdown: str
    images: list[EmbeddedImage] = field(default_factory=list)


# --------------------------------------------------------------------------
# 文本与图片工具
# --------------------------------------------------------------------------
def read_text_file(path: Path) -> str:
    """按 UTF-8 / GB18030 依次尝试读取文本。"""
    raw = path.read_bytes()[:MAX_TEXT_BYTES]
    return decode_text(raw)


def decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "big5"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _clip(text: str, limit: int = MAX_DOC_CHARS) -> str:
    if len(text) <= limit:
        return text
    return text[:limit] + f"\n\n> 资料过长，已截断到 {limit} 字符。"


def _normalize_suffix(value: str | None) -> str:
    suffix = (value or "").lower().strip()
    if suffix and not suffix.startswith("."):
        suffix = "." + suffix
    return suffix


def _guess_image_suffix(name: str = "", mime: str = "", data: bytes = b"") -> str:
    """返回可直接落盘的图片后缀；无法识别时返回空串（例如 EMF/WMF）。"""
    suffix = _normalize_suffix(Path(name).suffix)
    if suffix in IMAGE_SUFFIXES:
        return suffix
    mime = (mime or "").split(";")[0].strip().lower()
    if mime in _MIME_SUFFIX:
        return _MIME_SUFFIX[mime]
    head = data[:16]
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return ".png"
    if head.startswith(b"\xff\xd8\xff"):
        return ".jpg"
    if head.startswith(b"GIF8"):
        return ".gif"
    if head.startswith(b"BM"):
        return ".bmp"
    if head.startswith(b"RIFF") and b"WEBP" in head:
        return ".webp"
    if head.lstrip().startswith(b"<svg") or b"<svg" in head:
        return ".svg"
    return ""


def _decode_data_uri(uri: str) -> tuple[bytes, str] | None:
    match = _DATA_URI_RE.match(uri.strip())
    if not match:
        return None
    mime, encoding, payload = match.groups()
    try:
        if encoding:
            data = base64.b64decode(re.sub(r"\s+", "", payload))
        else:
            data = unquote(payload).encode("utf-8", errors="replace")
    except (ValueError, TypeError):
        return None
    return data, _guess_image_suffix(mime=mime, data=data)


class _Images:
    """单次转换内的图片收集器，负责去重并生成占位链接。"""

    def __init__(self, prefix: str = "") -> None:
        self.prefix = prefix
        self.items: list[EmbeddedImage] = []
        self._seen: dict[str, str] = {}

    def add(
        self,
        data: bytes | None,
        suffix: str = "",
        *,
        alt: str = "",
        title: str = "",
        mime: str = "",
    ) -> str | None:
        if not data:
            return None
        digest = hashlib.sha1(data).hexdigest()
        if digest in self._seen:
            return self._seen[digest]
        suffix = _guess_image_suffix(suffix, mime, data)
        if not suffix:
            return None
        key = f"{self.prefix}img{len(self.items) + 1}"
        self.items.append(
            EmbeddedImage(key=key, data=data, suffix=suffix, alt=alt, title=title or alt)
        )
        url = f"pptking-image://{key}"
        self._seen[digest] = url
        return url


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:g}"
    text = str(value).strip()
    return text.replace("\r\n", "<br>").replace("\n", "<br>").replace("|", "\\|")


def table_markdown(rows: list[list[object]], *, header: bool = True) -> str:
    """把二维数据转成 Markdown 表格，自动裁掉尾部空列。"""
    cleaned = [[_cell_text(cell) for cell in row] for row in rows if row]
    cleaned = [row for row in cleaned if any(cell for cell in row)]
    if not cleaned:
        return ""
    width = max(len(row) for row in cleaned)
    if width == 0:
        return ""
    for row in cleaned:
        row.extend([""] * (width - len(row)))
    while width > 1 and all(not row[width - 1] for row in cleaned):
        for row in cleaned:
            row.pop()
        width -= 1

    header_row = cleaned[0] if header else [""] * width
    body = cleaned[1:] if header else cleaned
    lines = [
        "| " + " | ".join(header_row) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


# --------------------------------------------------------------------------
# HTML -> Markdown
# --------------------------------------------------------------------------
_SKIP_TAGS = {"script", "style", "head", "noscript", "template", "iframe"}
_BLOCK_TAGS = {
    "p",
    "div",
    "section",
    "article",
    "header",
    "footer",
    "main",
    "aside",
    "figure",
    "figcaption",
    "form",
    "dl",
    "dt",
    "dd",
    "address",
}
ImageLoader = Callable[[str, str], str | None]


class _HtmlToMarkdown(HTMLParser):
    def __init__(self, image_loader: ImageLoader | None = None) -> None:
        super().__init__(convert_charrefs=True)
        self._loader = image_loader
        self._blocks: list[str] = []
        self._buf: list[str] = []
        self._skip = 0
        self._pre = 0
        self._pre_buf: list[str] = []
        self._heading = 0
        self._quote = 0
        self._lists: list[dict[str, object]] = []
        self._links: list[str | None] = []
        self._table: list[list[str]] | None = None
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    # -- 输出 -------------------------------------------------------------
    def _take(self) -> str:
        text = "".join(self._buf)
        self._buf.clear()
        text = re.sub(r"[ \t\r\f\v]+", " ", text)
        text = re.sub(r" *\n *", "\n", text)
        return text.strip()

    def _emit(self, text: str) -> None:
        if not text:
            return
        if self._quote > 0:
            text = "\n".join("> " + line for line in text.splitlines())
        self._blocks.append(text)

    def _flush(self) -> None:
        self._emit(self._take())

    @property
    def markdown(self) -> str:
        self._flush()
        text = "\n\n".join(block.strip() for block in self._blocks if block.strip())
        return re.sub(r"\n{3,}", "\n\n", text).strip()

    # -- HTMLParser -------------------------------------------------------
    def handle_starttag(self, tag: str, attrs) -> None:  # type: ignore[no-untyped-def]
        tag = tag.lower()
        values = {key.lower(): (value or "") for key, value in attrs}

        if self._skip:
            if tag in _SKIP_TAGS:
                self._skip += 1
            return
        if tag in _SKIP_TAGS:
            self._skip = 1
            return
        if self._pre:
            if tag == "pre":
                self._pre += 1
            return
        if self._table is not None:
            if tag == "tr":
                self._row = []
            elif tag in {"td", "th"}:
                self._cell = []
            elif tag == "br" and self._cell is not None:
                self._cell.append(" ")
            elif tag == "img":
                alt = values.get("alt") or values.get("title") or ""
                if self._cell is not None and alt:
                    self._cell.append(alt)
            return

        if tag == "table":
            self._flush()
            self._table = []
        elif re.fullmatch(r"h[1-6]", tag):
            self._flush()
            self._heading = int(tag[1])
        elif tag in _BLOCK_TAGS:
            self._flush()
        elif tag == "br":
            self._buf.append("\n")
        elif tag == "hr":
            self._flush()
            self._emit("---")
        elif tag in {"ul", "ol"}:
            self._flush()
            self._lists.append({"ordered": tag == "ol", "index": 0})
        elif tag == "li":
            self._flush()
            depth = max(0, len(self._lists) - 1)
            prefix = "- "
            if self._lists and self._lists[-1]["ordered"]:
                self._lists[-1]["index"] = int(self._lists[-1]["index"]) + 1
                prefix = f"{self._lists[-1]['index']}. "
            self._buf.append("  " * depth + prefix)
        elif tag == "pre":
            self._flush()
            self._pre = 1
            self._pre_buf = []
        elif tag == "code":
            self._buf.append("`")
        elif tag in {"strong", "b"}:
            self._buf.append("**")
        elif tag in {"em", "i"}:
            self._buf.append("*")
        elif tag == "blockquote":
            self._flush()
            self._quote += 1
        elif tag == "a":
            href = values.get("href", "").strip()
            if href and not href.startswith(("#", "javascript:")):
                self._buf.append("[")
                self._links.append(href)
            else:
                self._links.append(None)
        elif tag == "img":
            src = values.get("src", "").strip()
            alt = values.get("alt") or values.get("title") or ""
            url = self._loader(src, alt) if self._loader else None
            if url:
                self._buf.append(f"![{alt}]({url})")
            elif alt:
                self._buf.append(alt)
        elif tag == "sup":
            self._buf.append("^")
        elif tag == "sub":
            self._buf.append("_")

    def handle_startendtag(self, tag: str, attrs) -> None:  # type: ignore[no-untyped-def]
        self.handle_starttag(tag, attrs)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if self._skip:
            if tag in _SKIP_TAGS:
                self._skip -= 1
            return
        if self._pre:
            if tag == "pre":
                self._pre -= 1
                if self._pre == 0:
                    code = "".join(self._pre_buf).strip("\n")
                    self._pre_buf = []
                    self._emit(f"```\n{code}\n```")
            return
        if self._table is not None:
            if tag in {"td", "th"}:
                if self._cell is not None and self._row is not None:
                    self._row.append(" ".join(self._cell).strip())
                self._cell = None
            elif tag == "tr":
                if self._row:
                    self._table.append(self._row)
                self._row = None
            elif tag == "table":
                rendered = table_markdown(self._table)
                if rendered:
                    self._emit(rendered)
                self._table = None
                self._row = None
                self._cell = None
            return

        if re.fullmatch(r"h[1-6]", tag):
            text = self._take()
            if text:
                self._emit("#" * self._heading + " " + text)
            self._heading = 0
        elif tag in _BLOCK_TAGS or tag in {"tr", "td", "th"}:
            self._flush()
        elif tag == "li":
            self._flush()
        elif tag in {"ul", "ol"}:
            self._flush()
            if self._lists:
                self._lists.pop()
        elif tag == "code":
            self._buf.append("`")
        elif tag in {"strong", "b"}:
            self._buf.append("**")
        elif tag in {"em", "i"}:
            self._buf.append("*")
        elif tag == "blockquote":
            self._flush()
            self._quote = max(0, self._quote - 1)
        elif tag == "a":
            href = self._links.pop() if self._links else None
            if href:
                self._buf.append(f"]({href})")

    def handle_data(self, data: str) -> None:
        if self._skip:
            return
        if self._pre:
            self._pre_buf.append(data)
            return
        if self._table is not None:
            if self._cell is not None:
                self._cell.append(data)
            return
        self._buf.append(data)

    def close(self) -> None:
        super().close()
        self._flush()


def _html_to_markdown(html: str, image_loader: ImageLoader | None = None) -> str:
    parser = _HtmlToMarkdown(image_loader)
    parser.feed(html)
    parser.close()
    return parser.markdown


def _file_image_loader(base: Path, images: _Images) -> ImageLoader:
    def load(src: str, alt: str) -> str | None:
        if not src:
            return None
        if src.startswith("pptking-image://"):
            return src
        if src.startswith("data:"):
            decoded = _decode_data_uri(src)
            if not decoded:
                return None
            data, suffix = decoded
            return images.add(data, suffix, alt=alt)
        parsed = urlparse(src)
        if parsed.scheme and parsed.scheme not in {"file"}:
            return None
        raw = unquote(parsed.path if parsed.scheme else src)
        path = (base / raw).resolve()
        try:
            if not path.is_file() or path.suffix.lower() not in IMAGE_SUFFIXES:
                return None
            return images.add(path.read_bytes(), path.suffix, alt=alt, title=path.stem)
        except OSError:
            return None

    return load


def _zip_image_loader(zf: zipfile.ZipFile, base: str, images: _Images) -> ImageLoader:
    names = set(zf.namelist())

    def load(src: str, alt: str) -> str | None:
        if not src:
            return None
        if src.startswith("pptking-image://"):
            return src
        if src.startswith("data:"):
            decoded = _decode_data_uri(src)
            if not decoded:
                return None
            data, suffix = decoded
            return images.add(data, suffix, alt=alt)
        parsed = urlparse(src)
        if parsed.scheme and parsed.scheme not in {"file"}:
            return None
        raw = unquote(parsed.path if parsed.scheme else src)
        member = posixpath.normpath(posixpath.join(base, raw.lstrip("/")))
        if member not in names:
            return None
        try:
            data = zf.read(member)
        except (KeyError, OSError):
            return None
        return images.add(data, Path(member).suffix, alt=alt, title=Path(member).stem)

    return load


# --------------------------------------------------------------------------
# 各格式转换器
# --------------------------------------------------------------------------
def _convert_plain(path: Path, images: _Images, depth: int) -> str:
    text = read_text_file(path)
    if path.suffix.lower() in {".md", ".markdown"}:
        return _MD_DATA_IMAGE_RE.sub(_md_data_image_repl(images), text)
    return text


def _md_data_image_repl(images: _Images):
    def repl(match: re.Match[str]) -> str:
        alt = match.group(1)
        decoded = _decode_data_uri(match.group(2))
        if not decoded:
            return alt
        data, suffix = decoded
        url = images.add(data, suffix, alt=alt)
        return f"![{alt}]({url})" if url else alt

    return repl


def _convert_code(path: Path, images: _Images, depth: int) -> str:
    text = read_text_file(path)
    suffix = path.suffix.lower()
    if suffix == ".json":
        try:
            text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        except ValueError:
            pass
    language = _LANG.get(suffix, suffix.lstrip("."))
    return f"```{language}\n{text.strip()}\n```"


def _convert_csv(path: Path, images: _Images, depth: int) -> str:
    text = read_text_file(path)
    suffix = path.suffix.lower()
    delimiter = "\t" if suffix == ".tsv" else ","
    if suffix == ".csv":
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            pass
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows: list[list[str]] = []
    truncated = False
    for index, row in enumerate(reader):
        if index >= MAX_TABLE_ROWS:
            truncated = True
            break
        rows.append(row[:MAX_TABLE_COLS])
    table = table_markdown(rows)
    if truncated:
        table += f"\n\n> 仅展示前 {MAX_TABLE_ROWS} 行。"
    return table


def _convert_html(path: Path, images: _Images, depth: int) -> str:
    html = read_text_file(path)
    return _html_to_markdown(html, _file_image_loader(path.parent, images))


# -- DOCX ------------------------------------------------------------------
def _convert_docx(path: Path, images: _Images, depth: int) -> str:
    try:
        import mammoth
    except ImportError:
        return _convert_docx_fallback(path, images)

    def convert_image(image) -> dict[str, str]:
        with image.open() as stream:
            data = stream.read()
        url = images.add(data, mime=getattr(image, "content_type", ""), alt="")
        return {"src": url or ""}

    try:
        with path.open("rb") as stream:
            result = mammoth.convert_to_html(
                stream, convert_image=mammoth.images.img_element(convert_image)
            )
    except Exception:  # noqa: BLE001 - 回退到内置的轻量解析
        return _convert_docx_fallback(path, images)

    def loader(src: str, alt: str) -> str | None:
        if src.startswith("pptking-image://"):
            return src
        return None

    markdown = _html_to_markdown(result.value, loader)
    # mammoth 只识别标准 drawing 节点；个别生成器写出的图片会被漏掉，
    # 这里兜底把 word/media 里的图片也抽出来。
    if not images.items:
        try:
            with zipfile.ZipFile(path) as zf:
                extras = _docx_media_blocks(zf, images)
        except (OSError, zipfile.BadZipFile):
            extras = []
        if extras:
            markdown = (markdown + "\n\n" + "\n\n".join(extras)).strip()
    return markdown


def _docx_media_blocks(zf: zipfile.ZipFile, images: _Images) -> list[str]:
    blocks: list[str] = []
    for name in sorted(n for n in zf.namelist() if n.startswith("word/media/")):
        if Path(name).suffix.lower() not in IMAGE_SUFFIXES:
            continue
        try:
            data = zf.read(name)
        except (KeyError, OSError):
            continue
        url = images.add(data, Path(name).suffix, alt=Path(name).stem)
        if url:
            blocks.append(f"![{Path(name).stem}]({url})")
    return blocks


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _convert_docx_fallback(path: Path, images: _Images) -> str:
    """不依赖 mammoth 的基本 DOCX 解析（段落/标题/列表/表格/图片）。"""
    try:
        zf = zipfile.ZipFile(path)
    except (OSError, zipfile.BadZipFile) as exc:
        raise ConversionError("DOCX 文件损坏") from exc
    with zf:
        try:
            root = ET.fromstring(zf.read("word/document.xml"))
        except (KeyError, ET.ParseError) as exc:
            raise ConversionError("DOCX 缺少 word/document.xml") from exc
        rels = _docx_relationships(zf)
        body = root.find(f"{_W}body")
        if body is None:
            return ""
        blocks = _docx_blocks(body, rels, zf, images)
    return "\n\n".join(blocks)


def _docx_relationships(zf: zipfile.ZipFile) -> dict[str, str]:
    try:
        root = ET.fromstring(zf.read("word/_rels/document.xml.rels"))
    except (KeyError, ET.ParseError):
        return {}
    mapping: dict[str, str] = {}
    for rel in root:
        rel_id = rel.get("Id")
        target = rel.get("Target")
        if not rel_id or not target:
            continue
        mapping[rel_id] = posixpath.normpath(posixpath.join("word", target))
    return mapping


def _docx_blocks(body, rels, zf, images: _Images) -> list[str]:
    blocks: list[str] = []
    for child in body:
        if child.tag == f"{_W}p":
            text = _docx_paragraph(child, rels, zf, images)
            if not text:
                continue
            style = _docx_style(child)
            heading = re.search(r"(?:heading|标题)\s*([1-6])", style, re.I)
            if heading:
                blocks.append("#" * int(heading.group(1)) + " " + text)
            elif "quote" in style.lower() or "引用" in style:
                blocks.append("\n".join("> " + line for line in text.splitlines()))
            else:
                blocks.append(("- " if _docx_is_list(child) else "") + text)
        elif child.tag == f"{_W}tbl":
            rows = []
            for tr in child.findall(f"{_W}tr"):
                cells = []
                for tc in tr.findall(f"{_W}tc"):
                    parts = [
                        _docx_paragraph(p, rels, zf, images)
                        for p in tc.findall(f"{_W}p")
                    ]
                    cells.append(" ".join(part for part in parts if part))
                rows.append(cells)
            table = table_markdown(rows)
            if table:
                blocks.append(table)
        elif child.tag == f"{_W}sdt":
            content = child.find(f"{_W}sdtContent")
            if content is not None:
                blocks.extend(_docx_blocks(content, rels, zf, images))
    return blocks


def _docx_style(paragraph) -> str:
    ppr = paragraph.find(f"{_W}pPr")
    if ppr is None:
        return ""
    style = ppr.find(f"{_W}pStyle")
    return style.get(f"{_W}val", "") if style is not None else ""


def _docx_is_list(paragraph) -> bool:
    ppr = paragraph.find(f"{_W}pPr")
    return ppr is not None and ppr.find(f"{_W}numPr") is not None


def _docx_paragraph(paragraph, rels, zf, images: _Images) -> str:
    parts: list[str] = []
    for node in paragraph.iter():
        if node.tag == f"{_W}t":
            parts.append(node.text or "")
        elif node.tag == f"{_W}tab":
            parts.append("\t")
        elif node.tag in {f"{_W}br", f"{_W}cr"}:
            parts.append("\n")
        elif node.tag == f"{_A}blip":
            rel_id = node.get(f"{_R}embed")
            member = rels.get(rel_id or "")
            if not member:
                continue
            try:
                data = zf.read(member)
            except (KeyError, OSError):
                continue
            url = images.add(data, Path(member).suffix, alt=Path(member).stem)
            if url:
                parts.append(f"![{Path(member).stem}]({url})")
    return "".join(parts).strip()


# -- PPTX ------------------------------------------------------------------
def _convert_pptx(path: Path, images: _Images, depth: int) -> str:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        from pptx.util import Inches
    except ImportError as exc:
        raise ConversionError("解析 PPTX 需要 python-pptx") from exc

    try:
        presentation = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001 - python-pptx 会抛多种异常
        raise ConversionError("PPTX 文件损坏") from exc

    min_width = Inches(0.4)
    min_height = Inches(0.3)
    blocks: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        try:
            if slide.element.get("show") == "0":
                continue
        except AttributeError:
            pass
        title_shape = None
        try:
            title_shape = slide.shapes.title
        except (AttributeError, ValueError):
            title_shape = None
        title = (title_shape.text or "").strip() if title_shape is not None else ""
        blocks.append(f"## 第 {number} 页" + (f"：{title}" if title else ""))

        shapes = sorted(
            slide.shapes,
            key=lambda shape: (
                float("-inf") if shape.top is None else shape.top,
                float("-inf") if shape.left is None else shape.left,
            ),
        )
        blocks.extend(
            _pptx_shapes(
                shapes,
                title_shape,
                images,
                MSO_SHAPE_TYPE,
                min_width,
                min_height,
                title,
            )
        )

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            note_text = (notes.text or "").strip() if notes is not None else ""
            if note_text:
                blocks.append("\n".join("> " + line for line in note_text.splitlines()))
    return "\n\n".join(block for block in blocks if block.strip())


def _pptx_shapes(
    shapes,
    title_shape,
    images: _Images,
    shape_type,
    min_width,
    min_height,
    context_title: str = "",
) -> list[str]:
    blocks: list[str] = []
    for shape in shapes:
        if shape is title_shape:
            continue
        if shape.shape_type == shape_type.GROUP:
            nested = sorted(
                shape.shapes,
                key=lambda item: (
                    float("-inf") if item.top is None else item.top,
                    float("-inf") if item.left is None else item.left,
                ),
            )
            blocks.extend(
                _pptx_shapes(
                    nested,
                    title_shape,
                    images,
                    shape_type,
                    min_width,
                    min_height,
                    context_title,
                )
            )
            continue
        if getattr(shape, "has_table", False):
            rows = [
                [cell.text for cell in row.cells] for row in shape.table.rows
            ]
            table = table_markdown(rows)
            if table:
                blocks.append(table)
            continue
        if getattr(shape, "has_chart", False):
            blocks.append(_chart_markdown(shape.chart))
            continue
        if shape.shape_type == shape_type.PICTURE:
            width = shape.width or 0
            height = shape.height or 0
            if width < min_width or height < min_height:
                continue
            try:
                image = shape.image
                data = image.blob
            except (AttributeError, ValueError, OSError):
                continue
            alt = _pptx_alt(shape, context_title)
            url = images.add(data, image.ext, alt=alt, title=shape.name)
            if url:
                blocks.append(f"![{alt}]({url})")
            continue
        if getattr(shape, "has_text_frame", False):
            lines = _pptx_text_lines(shape.text_frame)
            if lines:
                blocks.append("\n".join(lines))
    return blocks


def _pptx_alt(shape, context_title: str = "") -> str:
    try:
        descr = shape._element.nvPicPr.cNvPr.get("descr")
    except AttributeError:
        descr = None
    if descr:
        return descr.strip()
    if context_title:
        return f"{context_title} 配图"
    return (shape.name or "幻灯片图片").strip()


def _pptx_text_lines(text_frame) -> list[str]:
    lines: list[str] = []
    for paragraph in text_frame.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        level = int(getattr(paragraph, "level", 0) or 0)
        bullet = _pptx_has_bullet(paragraph)
        lines.append("  " * level + ("- " if bullet else "") + text)
    return lines


def _pptx_has_bullet(paragraph) -> bool:
    try:
        from pptx.oxml.ns import qn

        ppr = paragraph._p.pPr
        if ppr is None:
            return False
        if ppr.find(qn("a:buNone")) is not None:
            return False
        return (
            ppr.find(qn("a:buChar")) is not None
            or ppr.find(qn("a:buAutoNum")) is not None
            or ppr.find(qn("a:buBlip")) is not None
        )
    except (AttributeError, ImportError, ValueError):
        return False


def _chart_markdown(chart) -> str:
    try:
        plot = chart.plots[0]
        categories = [str(category) for category in plot.categories]
        series = list(plot.series)
        rows: list[list[object]] = [["类别"] + [item.name or f"系列{i + 1}" for i, item in enumerate(series)]]
        for index, category in enumerate(categories):
            row: list[object] = [category]
            for item in series:
                try:
                    row.append(item.values[index])
                except (IndexError, TypeError):
                    row.append("")
            rows.append(row)
        return table_markdown(rows)
    except (AttributeError, IndexError, TypeError, ValueError):
        return ""


# -- Excel -----------------------------------------------------------------
def _convert_xlsx(path: Path, images: _Images, depth: int) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ConversionError("解析 XLSX 需要 openpyxl") from exc

    try:
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ConversionError("XLSX 文件损坏") from exc

    blocks: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[list[object]] = []
            truncated = False
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_TABLE_ROWS:
                    truncated = True
                    break
                values = list(row[:MAX_TABLE_COLS])
                if any(value not in (None, "") for value in values):
                    rows.append(values)
            table = table_markdown(rows)
            if not table:
                continue
            section = f"## 工作表：{sheet.title}\n\n{table}"
            if truncated:
                section += f"\n\n> 仅展示前 {MAX_TABLE_ROWS} 行。"
            blocks.append(section)
    finally:
        workbook.close()

    blocks.extend(_xlsx_media(path, images))
    return "\n\n".join(blocks)


def _xlsx_media(path: Path, images: _Images) -> list[str]:
    try:
        zf = zipfile.ZipFile(path)
    except (OSError, zipfile.BadZipFile):
        return []
    blocks: list[str] = []
    with zf:
        for name in sorted(n for n in zf.namelist() if n.startswith("xl/media/")):
            if Path(name).suffix.lower() not in IMAGE_SUFFIXES:
                continue
            try:
                data = zf.read(name)
            except (KeyError, OSError):
                continue
            url = images.add(data, Path(name).suffix, alt=Path(name).stem)
            if url:
                blocks.append(f"![{Path(name).stem}]({url})")
    return blocks


def _convert_xls(path: Path, images: _Images, depth: int) -> str:
    try:
        import xlrd
    except ImportError as exc:
        raise ConversionError("解析旧版 XLS 需要 xlrd") from exc
    try:
        book = xlrd.open_workbook(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ConversionError("XLS 文件损坏") from exc

    blocks: list[str] = []
    for sheet in book.sheets():
        rows: list[list[object]] = []
        for index in range(min(sheet.nrows, MAX_TABLE_ROWS)):
            rows.append([sheet.cell_value(index, col) for col in range(min(sheet.ncols, MAX_TABLE_COLS))])
        table = table_markdown(rows)
        if table:
            blocks.append(f"## 工作表：{sheet.name}\n\n{table}")
    return "\n\n".join(blocks)


# -- PDF -------------------------------------------------------------------
def _convert_pdf(path: Path, images: _Images, depth: int) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return _convert_pdf_pdfminer(path, images)

    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ConversionError("PDF 文件损坏") from exc

    blocks: list[str] = []
    total = len(reader.pages)
    text_pages = 0
    for index, page in enumerate(reader.pages[:MAX_PDF_PAGES], start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:  # noqa: BLE001
            text = ""
        blocks.append(f"## 第 {index} 页")
        if text:
            text_pages += 1
            blocks.append(text)
        for image in _pdf_images(page):
            try:
                data = image.data
                name = getattr(image, "name", "") or "pdf_image"
            except (AttributeError, OSError):
                continue
            alt = f"{path.stem} 第{index}页配图"
            url = images.add(data, Path(name).suffix, alt=alt)
            if url:
                blocks.append(f"![{alt}]({url})")
    if total > MAX_PDF_PAGES:
        blocks.append(f"> 仅解析前 {MAX_PDF_PAGES} 页（共 {total} 页）。")
    if text_pages == 0:
        blocks.append("> 未提取到文字，PDF 可能是扫描件，建议先做 OCR。")
    return "\n\n".join(blocks)


def _pdf_images(page):
    try:
        return list(page.images)
    except Exception:  # noqa: BLE001 - 不同 pypdf 版本 API 差异较大
        return []


def _convert_pdf_pdfminer(path: Path, images: _Images) -> str:
    try:
        from pdfminer.high_level import extract_text
    except ImportError as exc:
        raise ConversionError("解析 PDF 需要安装 pypdf 或 pdfminer.six") from exc
    try:
        text = extract_text(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ConversionError("PDF 文件损坏") from exc
    return text.strip()


# -- EPUB ------------------------------------------------------------------
def _epub_spine(zf: zipfile.ZipFile) -> list[str]:
    try:
        container = ET.fromstring(zf.read("META-INF/container.xml"))
        rootfile = next(
            (node for node in container.iter() if node.tag.endswith("rootfile")), None
        )
        opf_path = rootfile.get("full-path") if rootfile is not None else None
        if not opf_path:
            raise KeyError
        opf = ET.fromstring(zf.read(opf_path))
    except (KeyError, StopIteration, ET.ParseError):
        return []

    base = posixpath.dirname(opf_path)
    manifest: dict[str, str] = {}
    for node in opf.iter():
        if not node.tag.endswith("item"):
            continue
        item_id = node.get("id")
        href = node.get("href")
        media = (node.get("media-type") or "").lower()
        if item_id and href and ("html" in media or href.lower().endswith((".html", ".xhtml", ".htm"))):
            manifest[item_id] = posixpath.normpath(posixpath.join(base, href))
    spine: list[str] = []
    for node in opf.iter():
        if not node.tag.endswith("itemref"):
            continue
        idref = node.get("idref")
        if idref and idref in manifest:
            spine.append(manifest[idref])
    return spine


def _convert_epub(path: Path, images: _Images, depth: int) -> str:
    try:
        zf = zipfile.ZipFile(path)
    except (OSError, zipfile.BadZipFile) as exc:
        raise ConversionError("EPUB 文件损坏") from exc

    with zf:
        members = _epub_spine(zf)
        if not members:
            members = sorted(
                name
                for name in zf.namelist()
                if name.lower().endswith((".html", ".xhtml", ".htm"))
            )
        blocks: list[str] = []
        for index, member in enumerate(members[:MAX_ZIP_MEMBERS], start=1):
            try:
                raw = zf.read(member)
            except (KeyError, OSError):
                continue
            html = decode_text(raw)
            base = posixpath.dirname(member)
            chapter = _html_to_markdown(html, _zip_image_loader(zf, base, images))
            if chapter.strip():
                blocks.append(f"## 章节 {index}\n\n{chapter.strip()}")
    return "\n\n".join(blocks)


# -- ZIP -------------------------------------------------------------------
def _convert_zip(path: Path, images: _Images, depth: int) -> str:
    try:
        zf = zipfile.ZipFile(path)
    except (OSError, zipfile.BadZipFile) as exc:
        raise ConversionError("ZIP 文件损坏") from exc

    blocks: list[str] = []
    with zf, tempfile.TemporaryDirectory(prefix="ppt-king-zip-") as temp:
        temp_dir = Path(temp)
        infos = [
            info
            for info in zf.infolist()
            if not info.is_dir()
            and not info.filename.startswith("__MACOSX")
            and not Path(info.filename).name.startswith(".")
        ]
        for index, info in enumerate(infos[:MAX_ZIP_MEMBERS]):
            if info.file_size > 40_000_000:
                continue
            suffix = Path(info.filename).suffix.lower()
            if suffix in IMAGE_SUFFIXES:
                try:
                    data = zf.read(info)
                except (KeyError, OSError):
                    continue
                url = images.add(data, suffix, alt=Path(info.filename).stem, title=info.filename)
                if url:
                    blocks.append(f"![{Path(info.filename).stem}]({url})")
                continue
            if suffix not in SUPPORTED_SUFFIXES and suffix not in {".zip"}:
                continue
            if suffix == ".zip" and depth >= 2:
                continue
            try:
                data = zf.read(info)
            except (KeyError, OSError):
                continue
            target = temp_dir / f"member{index}{suffix}"
            try:
                target.write_bytes(data)
                child = convert_document(target, key_prefix=f"z{index}_", _depth=depth + 1)
            except ConversionError:
                continue
            except OSError:
                continue
            if child.markdown.strip():
                blocks.append(f"## 文件：{info.filename}\n\n{child.markdown.strip()}")
                images.items.extend(child.images)
    if not blocks:
        raise ConversionError("ZIP 中没有可读取的资料")
    return "\n\n".join(blocks)


# -- RTF / EML -------------------------------------------------------------
def _convert_rtf(path: Path, images: _Images, depth: int) -> str:
    raw = path.read_bytes().decode("latin-1", errors="replace")
    return _rtf_to_text(raw)


def _rtf_to_text(raw: str) -> str:
    raw = re.sub(r"\{\\\*.*?\}", "", raw, flags=re.S)
    raw = re.sub(
        r"\\'([0-9a-fA-F]{2})",
        lambda match: bytes([int(match.group(1), 16)]).decode("cp1252", errors="ignore"),
        raw,
    )
    raw = re.sub(r"\\u(-?\d+)\??", lambda match: chr(int(match.group(1)) % 65536), raw)
    raw = raw.replace("\\par", "\n").replace("\\line", "\n").replace("\\tab", "\t")
    raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", raw)
    raw = raw.replace("\\{", "{").replace("\\}", "}").replace("\\\\", "\\")
    raw = raw.replace("{", "").replace("}", "")
    return "\n".join(line.strip() for line in raw.splitlines() if line.strip())


def _convert_eml(path: Path, images: _Images, depth: int) -> str:
    try:
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    except Exception as exc:  # noqa: BLE001
        raise ConversionError("EML 文件损坏") from exc

    blocks: list[str] = []
    subject = str(message.get("Subject") or "").strip()
    if subject:
        blocks.append(f"# {subject}")
    meta = [
        f"- 发件人：{message.get('From', '')}",
        f"- 收件人：{message.get('To', '')}",
        f"- 时间：{message.get('Date', '')}",
    ]
    blocks.append("\n".join(meta))

    texts: list[str] = []
    htmls: list[str] = []
    for part in message.walk():
        if part.is_multipart():
            continue
        content_type = (part.get_content_type() or "").lower()
        if content_type.startswith("image/"):
            try:
                data = part.get_payload(decode=True) or b""
            except (ValueError, TypeError):
                continue
            url = images.add(data, mime=content_type, alt=part.get_filename() or "邮件图片")
            if url:
                blocks.append(f"![{part.get_filename() or '邮件图片'}]({url})")
            continue
        if content_type not in {"text/plain", "text/html"}:
            continue
        try:
            payload = part.get_payload(decode=True) or b""
            charset = part.get_content_charset() or "utf-8"
            text = payload.decode(charset, errors="replace")
        except (UnicodeError, LookupError):
            continue
        if content_type == "text/html":
            htmls.append(text)
        else:
            texts.append(text)

    if htmls:
        blocks.append(_html_to_markdown("\n".join(htmls), _file_image_loader(path.parent, images)))
    elif texts:
        blocks.append("\n".join(texts).strip())
    return "\n\n".join(block for block in blocks if block.strip())


# -- Notebook --------------------------------------------------------------
def _convert_ipynb(path: Path, images: _Images, depth: int) -> str:
    try:
        notebook = json.loads(read_text_file(path))
    except ValueError as exc:
        raise ConversionError("Notebook 不是合法 JSON") from exc

    language = "python"
    try:
        language = (
            notebook.get("metadata", {}).get("kernelspec", {}).get("language", "python")
            or "python"
        )
    except AttributeError:
        pass

    blocks: list[str] = []
    for cell in notebook.get("cells", []):
        kind = cell.get("cell_type")
        source = _join_notebook_source(cell.get("source"))
        if kind == "markdown" and source.strip():
            blocks.append(source.strip())
        elif kind == "code" and source.strip():
            blocks.append(f"```{language}\n{source.strip()}\n```")
            for output in cell.get("outputs", []) or []:
                blocks.extend(_notebook_output(output, images))
        elif kind == "raw" and source.strip():
            blocks.append(source.strip())
    return "\n\n".join(block for block in blocks if block.strip())


def _join_notebook_source(value) -> str:
    if isinstance(value, list):
        return "".join(str(part) for part in value)
    return str(value or "")


def _notebook_output(output: dict, images: _Images) -> list[str]:
    kind = output.get("output_type")
    blocks: list[str] = []
    if kind == "stream":
        text = _join_notebook_source(output.get("text"))
        if text.strip():
            blocks.append(f"```\n{text.rstrip()}\n```")
        return blocks
    if kind in {"execute_result", "display_data"}:
        data = output.get("data") or {}
        plain = _join_notebook_source(data.get("text/plain"))
        if plain.strip():
            blocks.append(f"```\n{plain.rstrip()}\n```")
        for mime in ("image/png", "image/jpeg", "image/gif"):
            payload = data.get(mime)
            if not payload:
                continue
            raw = _join_notebook_source(payload)
            try:
                decoded = base64.b64decode(raw)
            except (ValueError, TypeError):
                continue
            url = images.add(decoded, mime=mime, alt="Notebook 输出图")
            if url:
                blocks.append(f"![Notebook 输出图]({url})")
        return blocks
    if kind == "error":
        traceback = output.get("traceback") or []
        if traceback:
            blocks.append("```\n" + "\n".join(str(line) for line in traceback[-8:]) + "\n```")
        return blocks
    return blocks


# --------------------------------------------------------------------------
# 调度
# --------------------------------------------------------------------------
_CONVERTERS: dict[str, Callable[[Path, _Images, int], str]] = {
    **{suffix: _convert_plain for suffix in PLAIN_SUFFIXES},
    **{suffix: _convert_code for suffix in CODE_SUFFIXES},
    **{suffix: _convert_html for suffix in HTML_SUFFIXES},
    ".csv": _convert_csv,
    ".tsv": _convert_csv,
    ".docx": _convert_docx,
    ".pptx": _convert_pptx,
    ".xlsx": _convert_xlsx,
    ".xls": _convert_xls,
    ".pdf": _convert_pdf,
    ".epub": _convert_epub,
    ".zip": _convert_zip,
    ".ipynb": _convert_ipynb,
    ".rtf": _convert_rtf,
    ".eml": _convert_eml,
}


def is_supported(suffix: str) -> bool:
    return _normalize_suffix(suffix) in _CONVERTERS


def supported_suffixes() -> list[str]:
    return sorted(SUPPORTED_SUFFIXES | IMAGE_SUFFIXES)


def convert_document(
    path: str | Path, *, key_prefix: str = "", _depth: int = 0
) -> ConvertedDocument:
    """把单个文件转换成 :class:`ConvertedDocument`。"""
    source = Path(path)
    suffix = source.suffix.lower()
    converter = _CONVERTERS.get(suffix)
    if converter is None:
        raise ConversionError(f"暂不支持 {suffix or '无后缀'} 格式")
    images = _Images(key_prefix)
    try:
        markdown = converter(source, images, _depth)
    except ConversionError:
        raise
    except (OSError, zipfile.BadZipFile) as exc:
        raise ConversionError(f"读取失败：{exc}") from exc
    except Exception as exc:  # noqa: BLE001 - 第三方解析器异常种类繁多
        raise ConversionError(f"{type(exc).__name__}: {exc}") from exc

    if len(markdown) > MAX_DOC_CHARS:
        markdown = _clip(markdown)
    if not markdown.strip() and not images.items:
        raise ConversionError("未提取到可读内容")
    return ConvertedDocument(markdown=markdown, images=images.items)
