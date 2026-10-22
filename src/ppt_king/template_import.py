"""从 PPTX 近似反推模板 Markdown。

不调用模型：按形状包围盒映射到 16×9 网格，识别文字、图片和表格，
每页生成一个模板文件。适合把手工排好的模板 PPT 快速转回可编辑模板。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

from .grid import GRID_COLS, GRID_ROWS

_SAFE_NAME_RE = re.compile(r"[^\w\u4e00-\u9fff.-]+", re.UNICODE)
_RED = "FF0000"


class TemplateImportError(ValueError):
    """PPTX 无法转换为模板。"""


@dataclass
class ImportedTemplate:
    title: str
    markdown: str
    slide_index: int = 0
    path: Path | None = None
    images: list[Path] = field(default_factory=list)


def _safe_name(value: str, fallback: str) -> str:
    name = _SAFE_NAME_RE.sub("_", value).strip("._ ")
    return name[:40] or fallback


def _grid_box(
    left: int,
    top: int,
    width: int,
    height: int,
    slide_w: int,
    slide_h: int,
) -> tuple[int, int, int, int]:
    cell_w = slide_w / GRID_COLS
    cell_h = slide_h / GRID_ROWS

    def x_of(value: float) -> int:
        return max(1, min(GRID_COLS, round(value / cell_w) + 1))

    def y_of(value: float) -> int:
        return max(1, min(GRID_ROWS, round(value / cell_h) + 1))

    def x_end(value: float) -> int:
        return max(1, min(GRID_COLS, round(value / cell_w)))

    def y_end(value: float) -> int:
        return max(1, min(GRID_ROWS, round(value / cell_h)))

    col1 = x_of(left)
    row1 = y_of(top)
    col2 = max(col1, x_end(left + width))
    row2 = max(row1, y_end(top + height))
    return col1, row1, col2, row2


def _run_markdown(run) -> str:
    text = run.text
    if not text:
        return ""
    red = False
    try:
        red = run.font.color.type is not None and str(run.font.color.rgb) == _RED
    except (AttributeError, ValueError):
        red = False
    if red:
        return f"<red>{text}</red>"
    if run.font.bold:
        return f"**{text}**"
    return text


def _paragraph_markdown(paragraph) -> str:
    text = "".join(_run_markdown(run) for run in paragraph.runs).strip()
    if not text:
        return ""
    if text.startswith("• "):
        text = "- " + text[2:]
    elif text.startswith("– "):
        text = "  - " + text[2:]
    return text


def _text_markdown(text_frame) -> tuple[str, str, float | None]:
    lines: list[str] = []
    align = ""
    size: float | None = None
    for paragraph in text_frame.paragraphs:
        line = _paragraph_markdown(paragraph)
        if line:
            lines.append(line)
        for run in paragraph.runs:
            if run.font.size and (size is None or run.font.size.pt > size):
                size = run.font.size.pt
        if paragraph.alignment == PP_ALIGN.CENTER:
            align = "center"
        elif paragraph.alignment == PP_ALIGN.RIGHT and not align:
            align = "right"
    return "\n".join(lines), align, size


def _cell_text(cell) -> str:
    text = " ".join(
        paragraph.text.strip()
        for paragraph in cell.text_frame.paragraphs
        if paragraph.text.strip()
    )
    return text.replace("|", "\\|").replace("\n", " ")


def _table_markdown(table) -> str:
    rows = list(table.rows)
    if not rows:
        return ""
    header = [_cell_text(cell) for cell in rows[0].cells]
    if not header:
        return ""
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in rows[1:]:
        cells = [_cell_text(cell) for cell in row.cells]
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def _image_alt(shape) -> str:
    try:
        value = shape._element.nvPicPr.cNvPr.get("descr") or ""
    except AttributeError:
        value = ""
    return value.strip() or "配图"


def _is_logo(shape, slide_w: int, slide_h: int) -> bool:
    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
        return False
    cell_w = slide_w / GRID_COLS
    cell_h = slide_h / GRID_ROWS
    return (
        shape.left <= cell_w * 0.2
        and shape.top <= cell_h * 0.2
        and shape.width <= cell_w * 4.5
        and shape.height <= cell_h * 1.2
    )


def _is_divider(shape, slide_w: int) -> bool:
    return (
        shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
        and shape.height <= 15000
        and shape.width >= slide_w * 0.5
    )


def _find_title_shape(shapes, slide_w: int, slide_h: int):
    cell_h = slide_h / GRID_ROWS
    candidates = []
    for shape in shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = shape.text_frame.text.strip()
        if not text:
            continue
        top_units = shape.top / cell_h
        width_units = shape.width / (slide_w / GRID_COLS)
        if 0.6 <= top_units <= 1.8 and width_units >= 10:
            candidates.append((shape.top, shape.left, text, shape))
    if not candidates:
        return None, ""
    candidates.sort(key=lambda item: (item[0], item[1]))
    return candidates[0][3], candidates[0][2]


def _block_anchor(
    col1: int,
    row1: int,
    col2: int,
    row2: int,
    align: str = "",
    size: float | None = None,
) -> str:
    parts = [f"@{col1},{row1}-{col2},{row2}"]
    if align:
        parts.append(align)
    if size:
        parts.append(f"{round(size):g}")
    return " ".join(parts)


def extract_templates(pptx_path: str | Path) -> list[ImportedTemplate]:
    """从 PPTX 提取每页模板；图片先保留在内存中的引用描述。"""
    source = Path(pptx_path)
    if not source.is_file():
        raise TemplateImportError(f"找不到 PPTX：{source}")
    prs = Presentation(str(source))
    slide_w, slide_h = prs.slide_width, prs.slide_height
    templates: list[ImportedTemplate] = []

    for slide_no, slide in enumerate(prs.slides, 1):
        shapes = [
            shape
            for shape in slide.shapes
            if None not in (shape.left, shape.top, shape.width, shape.height)
        ]
        title_shape, title_text = _find_title_shape(shapes, slide_w, slide_h)
        if not title_text:
            try:
                title_text = slide.shapes.title.text.strip()
            except (AttributeError, ValueError):
                title_text = ""
        title = title_text or f"slide-{slide_no:02d}"
        blocks: list[tuple[int, int, str]] = []
        images: list[Path] = []

        for shape in shapes:
            if shape is title_shape:
                continue
            if _is_divider(shape, slide_w) or _is_logo(shape, slide_w, slide_h):
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                continue

            col1, row1, col2, row2 = _grid_box(
                shape.left,
                shape.top,
                shape.width,
                shape.height,
                slide_w,
                slide_h,
            )

            if shape.has_table:
                text = _table_markdown(shape.table)
                if text:
                    blocks.append((row1, col1, f"{_block_anchor(col1, row1, col2, row2)}\n{text}"))
                continue

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                alt = _image_alt(shape)
                ext = (shape.image.ext or "png").lstrip(".")
                filename = f"{_safe_name(title, f'slide-{slide_no:02d}')}_img{len(images) + 1:02d}.{ext}"
                images.append(Path(f"assets/{filename}"))
                blocks.append(
                    (
                        row1,
                        col1,
                        f"{_block_anchor(col1, row1, col2, row2)}\n![{alt}](assets/{filename})",
                    )
                )
                continue

            if shape.has_text_frame:
                text, align, size = _text_markdown(shape.text_frame)
                if text:
                    blocks.append(
                        (
                            row1,
                            col1,
                            f"{_block_anchor(col1, row1, col2, row2, align, size)}\n{text}",
                        )
                    )

        if not blocks:
            continue
        blocks.sort(key=lambda item: (item[0], item[1]))
        markdown = "\n\n".join(block for _, _, block in blocks).strip() + "\n"
        templates.append(
            ImportedTemplate(
                title=title,
                markdown=markdown,
                slide_index=slide_no - 1,
                images=images,
            )
        )

    if not templates:
        raise TemplateImportError(f"{source.name} 中没有可提取的内容")
    return templates


def write_templates(
    pptx_path: str | Path,
    out_dir: str | Path,
    *,
    force: bool = False,
) -> list[ImportedTemplate]:
    """从 PPTX 生成模板文件；默认不覆盖已有文件。"""
    source = Path(pptx_path)
    target = Path(out_dir)
    target.mkdir(parents=True, exist_ok=True)
    assets = target / "assets"
    assets.mkdir(exist_ok=True)
    prs = Presentation(str(source))
    slide_w, slide_h = prs.slide_width, prs.slide_height
    slides = list(prs.slides)
    results = extract_templates(source)
    used: dict[str, int] = {}

    for template in results:
        slide = slides[template.slide_index]
        stem = _safe_name(
            template.title, f"slide-{template.slide_index + 1:02d}"
        )
        count = used.get(stem, 0) + 1
        used[stem] = count
        if count > 1:
            stem = f"{stem}-{count}"
        path = target / f"{stem}.md"
        if path.exists() and not force:
            continue

        shapes = [
            shape
            for shape in slide.shapes
            if None not in (shape.left, shape.top, shape.width, shape.height)
        ]
        title_shape, _ = _find_title_shape(shapes, slide_w, slide_h)
        image_no = 0
        for shape in shapes:
            if shape is title_shape or _is_logo(shape, slide_w, slide_h):
                continue
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            image_no += 1
            ext = (shape.image.ext or "png").lstrip(".")
            filename = f"{stem}_img{image_no:02d}.{ext}"
            (assets / filename).write_bytes(shape.image.blob)
            template.images.append(assets / filename)
        template.path = path
        path.write_text(template.markdown, encoding="utf-8")
    return [item for item in results if item.path is not None]
