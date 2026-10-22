"""GridDeck -> PPTX：16×9 坐标定位的确定性渲染器。

风格目标与示例一致：白底、左对齐加粗标题、正文单栏、蓝色表头、无装饰图形。
所有几何来自格子坐标，字号只做「按框宽起步、按框高下调」的确定性计算。
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass, field, replace
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

DEFAULT_LOGO = Path(__file__).resolve().parents[2] / "static" / "page_logo.png"


from .grid import (
    CODE,
    EMU_PER_CELL,
    FORMULA,
    GRID_COLS,
    GRID_ROWS,
    IMAGE,
    QUOTE,
    TABLE,
    TEXT,
    Box,
    FlowItem,
    GridBlock,
    GridDeck,
    GridSlide,
)

SLIDE_W = EMU_PER_CELL * GRID_COLS  # 12192000
SLIDE_H = EMU_PER_CELL * GRID_ROWS  # 6858000

EMU_PER_PT = 12700
PAD = 72000  # 块内边距 0.08 英寸
TITLE_ROWS = 1
LOGO_COLS = 8
LOGO_ROWS = 1
BOTTOM_SAFE = EMU_PER_CELL  # 最底下一行留空，避免正文贴边
COVER_TITLES = {"封面", "封面页", "标题页", "cover", "title"}


@dataclass(frozen=True)
class GridTheme:
    font: str = "Times New Roman"
    background: RGBColor = RGBColor(0xFF, 0xFF, 0xFF)
    title: RGBColor = RGBColor(0x00, 0x00, 0x00)
    body: RGBColor = RGBColor(0x00, 0x00, 0x00)
    muted: RGBColor = RGBColor(0x59, 0x59, 0x59)
    danger: RGBColor = RGBColor(0xFF, 0x00, 0x00)
    accent: RGBColor = RGBColor(0xC0, 0x00, 0x00)
    table_head: RGBColor = RGBColor(0x2E, 0x5C, 0x8A)
    table_head_text: RGBColor = RGBColor(0xFF, 0xFF, 0xFF)
    table_alt: RGBColor = RGBColor(0xF2, 0xF5, 0xF9)
    rule: RGBColor = RGBColor(0xD0, 0xD7, 0xE0)

    title_pt: float = 28
    body_pt: float = 20
    body_max_pt: float = 24
    min_pt: float = 15
    caption_pt: float = 18
    table_pt: float = 20
    line_spacing: float = 1.45


DEFAULT_THEME = GridTheme()


@dataclass
class RenderReport:
    slides: int = 0
    blocks: int = 0
    auto_blocks: int = 0
    overflows: list[str] = field(default_factory=list)
    missing_images: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------
# 度量
# --------------------------------------------------------------------------
def _units(text: str) -> float:
    total = 0.0
    for char in text:
        if char in " \t":
            total += 0.30
        elif ord(char) < 0x2E80:
            total += 0.54
        else:
            total += 1.0
    return total or 1.0


def _lines(text: str, width_emu: int, size_pt: float) -> int:
    usable = max(1.0, (width_emu - 2 * PAD) / EMU_PER_PT)
    per_line = max(1.0, usable / max(1.0, size_pt))
    return max(1, math.ceil(_units(text) / per_line))


def _line_h(size_pt: float, spacing: float) -> int:
    return int(size_pt * spacing * EMU_PER_PT)


def _text_h(text: str, width_emu: int, size_pt: float, spacing: float = 1.45) -> int:
    return _lines(text, width_emu, size_pt) * _line_h(size_pt, spacing) + 2 * PAD


def _flow_h(
    flow: list[FlowItem], width: int, size_pt: float, spacing: float = 1.45
) -> int:
    total = 2 * PAD
    for item in flow:
        indent = int(0.32 * 914400 * item.level)
        inner = max(1, width - indent)
        delta = 2 if item.kind == "heading" else 0
        total += _lines(item.text, inner, size_pt + delta) * _line_h(size_pt + delta, spacing)
        total += int(0.10 * 914400)
    return total


def _aspect(path: Path) -> float:
    try:
        with Image.open(path) as image:
            width, height = image.size
        if width and height:
            return width / height
    except Exception:
        pass
    return 4 / 3


# --------------------------------------------------------------------------
# 基础绘制
# --------------------------------------------------------------------------
def _style(
    run,
    size_pt: float,
    theme: GridTheme,
    *,
    bold=False,
    italic=False,
    color=None,
    mono=False,
) -> None:
    font = run.font
    font.size = Pt(size_pt)
    font.bold = bold
    font.italic = italic
    font.name = "Consolas" if mono else theme.font
    font.color.rgb = color or theme.body
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rpr, qn("a:ea"))
    ea.set("typeface", "Consolas" if mono else theme.font)


def _frame(slide, x: int, y: int, w: int, h: int, anchor=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = Emu(PAD)
    frame.margin_right = Emu(PAD)
    frame.margin_top = Emu(PAD // 2)
    frame.margin_bottom = Emu(PAD // 2)
    return frame


def _align(value: str) -> PP_ALIGN:
    return {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(value, PP_ALIGN.LEFT)


@dataclass(frozen=True)
class InlineRun:
    text: str
    bold: bool = False
    red: bool = False


_INLINE_RE = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|<red>.*?</red>)", re.IGNORECASE
)


def _segments(
    text: str, *, bold: bool = False, red: bool = False
) -> list[InlineRun]:
    """拆分行内格式：**加粗**、<red>红色</red>；反引号只去掉不显示。"""
    parts: list[InlineRun] = []
    last = 0
    for match in _INLINE_RE.finditer(text):
        if match.start() > last:
            parts.append(InlineRun(text[last : match.start()], bold, red))
        chunk = match.group(0)
        if chunk.startswith("**") and chunk.endswith("**"):
            parts.extend(_segments(chunk[2:-2], bold=True, red=red))
        elif chunk.startswith("`") and chunk.endswith("`"):
            parts.append(InlineRun(chunk[1:-1], bold, red))
        else:
            inner = re.sub(r"</?red>", "", chunk, flags=re.IGNORECASE)
            parts.extend(_segments(inner, bold=bold, red=True))
        last = match.end()
    if last < len(text):
        parts.append(InlineRun(text[last:], bold, red))
    return [part for part in parts if part.text]


# --------------------------------------------------------------------------
# 文本流
# --------------------------------------------------------------------------
def _draw_flow(
    slide, block: GridBlock, box: Box, theme: GridTheme, report: RenderReport, tag: str
) -> None:
    def fits(size: float) -> bool:
        return _flow_h(block.flow, box.width, size, theme.line_spacing) <= box.height

    if box.size is not None:
        size = max(box.size, theme.min_pt)
        while size > theme.min_pt and not fits(size):
            size -= 1
    else:
        # 先从基准字号向上撑满，再向下收缩
        size = max(theme.body_pt, theme.min_pt)
        while size < theme.body_max_pt and fits(size + 1):
            size += 1
        while size > theme.min_pt and not fits(size):
            size -= 1

    if not fits(size):
        report.overflows.append(f"{tag}「{block.preview[:14]}」")

    anchor = MSO_ANCHOR.MIDDLE if box.align in ("center", "right") else MSO_ANCHOR.TOP
    frame = _frame(slide, box.x, box.y, box.width, box.height, anchor=anchor)
    counters = [0, 0, 0, 0]
    for index, item in enumerate(block.flow):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = _align(box.align)
        paragraph.line_spacing = theme.line_spacing
        paragraph.space_before = Pt(5 if index else 0)

        level = min(item.level, 3)
        if item.kind == "bullet":
            text = ("• " if level == 0 else "– ") + item.text
            run_size = max(theme.min_pt, size - 1.5 * level)
        elif item.kind == "ordered":
            counters[level] += 1
            for deeper in range(level + 1, 4):
                counters[deeper] = 0
            text = f"{counters[level]}. {item.text}"
            run_size = max(theme.min_pt, size - 1.5 * level)
        elif item.kind == "heading":
            text = item.text
            run_size = size + 2
        else:
            text = item.text
            run_size = size

        if level and item.kind in ("bullet", "ordered"):
            paragraph.level = level

        for segment in _segments(text, bold=item.kind == "heading"):
            run = paragraph.add_run()
            run.text = segment.text
            _style(
                run,
                run_size,
                theme,
                bold=segment.bold or segment.red,
                color=theme.danger if segment.red else None,
            )


def _draw_quote(
    slide, block: GridBlock, box: Box, theme: GridTheme, report: RenderReport, tag: str
) -> None:
    text = block.flow[0].text if block.flow else ""
    bar = int(0.05 * 914400)
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, int(box.x + PAD), int(box.y + PAD), bar, int(box.height - 2 * PAD)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = theme.rule
    shape.line.fill.background()
    shape.shadow.inherit = False

    size = max(theme.min_pt, theme.body_pt - 1)
    width = box.width - bar * 3
    while (
        size > theme.min_pt
        and _text_h(text, width, size, theme.line_spacing) > box.height
    ):
        size -= 1
    frame = _frame(slide, box.x + bar * 3, box.y, width, box.height)
    paragraph = frame.paragraphs[0]
    paragraph.line_spacing = theme.line_spacing
    for segment in _segments(text):
        run = paragraph.add_run()
        run.text = segment.text
        _style(
            run,
            size,
            theme,
            bold=segment.bold or segment.red,
            color=theme.danger if segment.red else theme.body,
        )


_FORMULA_SYMBOLS = {
    "alpha": "α",
    "beta": "β",
    "gamma": "γ",
    "delta": "δ",
    "epsilon": "ε",
    "varepsilon": "ε",
    "zeta": "ζ",
    "eta": "η",
    "theta": "θ",
    "iota": "ι",
    "kappa": "κ",
    "lambda": "λ",
    "mu": "μ",
    "nu": "ν",
    "xi": "ξ",
    "pi": "π",
    "rho": "ρ",
    "sigma": "σ",
    "tau": "τ",
    "upsilon": "υ",
    "phi": "φ",
    "varphi": "φ",
    "chi": "χ",
    "psi": "ψ",
    "omega": "ω",
    "Gamma": "Γ",
    "Delta": "Δ",
    "Theta": "Θ",
    "Lambda": "Λ",
    "Xi": "Ξ",
    "Pi": "Π",
    "Sigma": "Σ",
    "Phi": "Φ",
    "Psi": "Ψ",
    "Omega": "Ω",
    "le": "≤",
    "leq": "≤",
    "ge": "≥",
    "geq": "≥",
    "ne": "≠",
    "neq": "≠",
    "approx": "≈",
    "equiv": "≡",
    "sim": "∼",
    "propto": "∝",
    "times": "×",
    "cdot": "·",
    "pm": "±",
    "mp": "∓",
    "div": "÷",
    "rightarrow": "→",
    "to": "→",
    "leftarrow": "←",
    "leftrightarrow": "↔",
    "Rightarrow": "⇒",
    "Leftarrow": "⇐",
    "Leftrightarrow": "⇔",
    "infty": "∞",
    "partial": "∂",
    "nabla": "∇",
    "sum": "Σ",
    "prod": "Π",
    "int": "∫",
    "in": "∈",
    "notin": "∉",
    "subset": "⊂",
    "subseteq": "⊆",
    "cup": "∪",
    "cap": "∩",
    "forall": "∀",
    "exists": "∃",
    "ldots": "…",
    "cdots": "⋯",
    "degree": "°",
    "quad": " ",
    "qquad": "  ",
}
_FORMULA_WRAPPER_RE = re.compile(
    r"\\(?:text|mathrm|mathit|mathbf|mathsf|mathtt|operatorname|vec|hat|bar|overline|underline)"
    r"\{([^{}]*)\}"
)
_FORMULA_FRAC_RE = re.compile(r"\\frac\{([^{}]*)\}\{([^{}]*)\}")
_FORMULA_SQRT_RE = re.compile(r"\\sqrt\{([^{}]*)\}")
_FORMULA_COMMAND_RE = re.compile(r"\\([A-Za-z]+)")


def _normalize_formula(text: str) -> str:
    """把常见 LaTeX 公式命令转成 PowerPoint 可直接显示的 Unicode 文本。"""
    text = text.strip()
    if text.startswith("$$"):
        text = text[2:]
    if text.endswith("$$"):
        text = text[:-2]
    text = text.strip()
    if not text:
        return text

    previous = None
    while previous != text:
        previous = text
        text = _FORMULA_WRAPPER_RE.sub(r"\1", text)
    previous = None
    while previous != text:
        previous = text
        text = _FORMULA_FRAC_RE.sub(r"(\1)/(\2)", text)
    text = _FORMULA_SQRT_RE.sub(r"√(\1)", text)
    text = text.replace("\\left", "").replace("\\right", "")
    text = (
        text.replace("\\{", "{")
        .replace("\\}", "}")
        .replace("\\_", "_")
        .replace("\\%", "%")
        .replace("\\&", "&")
        .replace("\\#", "#")
        .replace("\\,", " ")
        .replace("\\;", " ")
        .replace("\\:", " ")
        .replace("\\!", "")
        .replace("\\ ", " ")
    )
    text = _FORMULA_COMMAND_RE.sub(
        lambda match: _FORMULA_SYMBOLS.get(match.group(1), ""), text
    )
    text = text.replace("{", "").replace("}", "").replace("~", " ")
    return re.sub(r"[ \t]+", " ", text).strip() or text


def _draw_formula(
    slide, block: GridBlock, box: Box, theme: GridTheme, report: RenderReport, tag: str
) -> None:
    text = _normalize_formula(block.flow[0].text if block.flow else "")
    size = box.size or 24
    while size > theme.min_pt and _text_h(text, box.width, size, 1.2) > box.height:
        size -= 1
    if _text_h(text, box.width, size, 1.2) > box.height:
        report.overflows.append(f"{tag}「{block.preview[:14]}」")
    frame = _frame(
        slide, box.x, box.y, box.width, box.height, anchor=MSO_ANCHOR.MIDDLE
    )
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.line_spacing = 1.2
    run = paragraph.add_run()
    run.text = text
    _style(run, size, theme, italic=True, color=theme.body)


def _draw_code(
    slide, block: GridBlock, box: Box, theme: GridTheme, report: RenderReport, tag: str
) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, int(box.x), int(box.y), int(box.width), int(box.height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xF6, 0xF8, 0xFA)
    shape.line.fill.background()
    shape.shadow.inherit = False

    text = block.flow[0].text if block.flow else ""
    size = max(theme.min_pt, theme.body_pt - 3)
    while size > theme.min_pt and _text_h(text, box.width, size, 1.3) > box.height:
        size -= 1
    frame = _frame(slide, box.x, box.y, box.width, box.height)
    paragraph = frame.paragraphs[0]
    paragraph.line_spacing = 1.3
    run = paragraph.add_run()
    run.text = text
    _style(run, size, theme, mono=True)


# --------------------------------------------------------------------------
# 图片 / 表格
# --------------------------------------------------------------------------
def _draw_image(
    slide,
    block: GridBlock,
    box: Box,
    theme: GridTheme,
    root: Path,
    report: RenderReport,
    tag: str,
) -> None:
    path = (root / block.path).resolve() if block.path else None
    if path is None or not path.is_file():
        frame = _frame(slide, box.x, box.y, box.width, box.height)
        paragraph = frame.paragraphs[0]
        run = paragraph.add_run()
        run.text = f"[图片缺失：{block.path}]"
        _style(run, theme.caption_pt + 2, theme, color=theme.muted)
        report.missing_images.append(f"{tag} {block.path or '(空路径)'}")
        return

    caption_h = int(theme.caption_pt * 1.8 * EMU_PER_PT) + PAD if block.caption else 0
    avail_h = max(int(0.4 * 914400), box.height - caption_h)
    avail_w = max(int(0.4 * 914400), box.width - 2 * PAD)
    aspect = _aspect(path)
    draw_w, draw_h = avail_w, avail_w / aspect
    if draw_h > avail_h:
        draw_h, draw_w = avail_h, avail_h * aspect
    x = box.x + (box.width - draw_w) / 2
    y = box.y + (avail_h - draw_h) / 2
    slide.shapes.add_picture(str(path), int(x), int(y), int(draw_w), int(draw_h))

    if block.caption:
        frame = _frame(slide, box.x, int(y + draw_h), box.width, caption_h)
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = block.caption
        _style(run, theme.caption_pt, theme, color=theme.body)


def _draw_table(
    slide, block: GridBlock, box: Box, theme: GridTheme, report: RenderReport, tag: str
) -> None:
    rows = 1 + len(block.rows)
    cols = max(1, len(block.header))
    table_w = box.width - 2 * PAD
    avail_h = max(int(0.5 * 914400), box.height - 2 * PAD)

    weights: list[float] = []
    for index in range(cols):
        longest = _units(block.header[index]) if index < len(block.header) else 1.0
        for row in block.rows:
            if index < len(row):
                longest = max(longest, _units(row[index]))
        weights.append(max(3.0, min(longest, 22.0)))
    total = sum(weights) or 1.0
    col_widths = [max(1, int(table_w * weight / total)) for weight in weights]

    def row_heights(size: float) -> list[int]:
        heights: list[int] = []
        for row_index in range(rows):
            data = block.header if row_index == 0 else block.rows[row_index - 1]
            max_lines = 1
            for col_index in range(cols):
                value = data[col_index] if col_index < len(data) else ""
                max_lines = max(max_lines, _lines(value, col_widths[col_index], size))
            line_h = int(size * 1.2 * EMU_PER_PT)
            heights.append(
                max(int(0.32 * 914400), max_lines * line_h + 2 * (PAD // 3))
            )
        return heights

    size = theme.table_pt
    heights = row_heights(size)
    while size > theme.min_pt and sum(heights) > avail_h:
        size -= 1
        heights = row_heights(size)
    if sum(heights) > avail_h:
        report.overflows.append(
            f"{tag} 表格 {rows} 行，{size:.0f}pt 仍装不下"
        )
        scale = avail_h / max(1, sum(heights))
        heights = [max(int(0.24 * 914400), int(height * scale)) for height in heights]

    table_h = sum(heights)
    shape = slide.shapes.add_table(
        rows,
        cols,
        int(box.x + PAD),
        int(box.y + PAD),
        int(table_w),
        int(table_h),
    )
    table = shape.table
    table.first_row = True
    table.horz_banding = False

    for index, width in enumerate(col_widths):
        table.columns[index].width = Emu(width)
    for index, height in enumerate(heights):
        table.rows[index].height = Emu(height)

    for row_index in range(rows):
        data = block.header if row_index == 0 else block.rows[row_index - 1]
        for col_index in range(cols):
            cell = table.cell(row_index, col_index)
            cell.margin_left = Emu(PAD)
            cell.margin_right = Emu(PAD)
            cell.margin_top = Emu(PAD // 3)
            cell.margin_bottom = Emu(PAD // 3)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if row_index == 0:
                cell.fill.fore_color.rgb = theme.table_head
            else:
                cell.fill.fore_color.rgb = (
                    theme.table_alt if row_index % 2 == 0 else theme.background
                )
            frame = cell.text_frame
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.LEFT if col_index == 0 else PP_ALIGN.CENTER
            value = data[col_index] if col_index < len(data) else ""
            for segment in _segments(value, bold=row_index == 0):
                run = paragraph.add_run()
                run.text = segment.text
                _style(
                    run,
                    size,
                    theme,
                    bold=segment.bold or segment.red,
                    color=(
                        theme.table_head_text
                        if row_index == 0
                        else (theme.danger if segment.red else theme.body)
                    ),
                )


# --------------------------------------------------------------------------
# 页级
# --------------------------------------------------------------------------
def _draw_header_text(slide, text: str, theme: GridTheme) -> None:
    frame = _frame(
        slide,
        8 * EMU_PER_CELL,
        0,
        8 * EMU_PER_CELL,
        LOGO_ROWS * EMU_PER_CELL,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.RIGHT
    run = paragraph.add_run()
    run.text = text
    _style(run, theme.body_pt, theme, color=theme.accent)


def _draw_header_rule(slide, theme: GridTheme) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0,
        LOGO_ROWS * EMU_PER_CELL - 9000,
        SLIDE_W,
        9000,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = theme.rule
    shape.line.fill.background()
    shape.shadow.inherit = False


def _draw_logo(slide, path: Path) -> None:
    width = LOGO_COLS * EMU_PER_CELL
    height = LOGO_ROWS * EMU_PER_CELL
    aspect = _aspect(path)
    draw_w, draw_h = width, width / aspect
    if draw_h > height:
        draw_h, draw_w = height, height * aspect
    x = 0
    y = (height - draw_h) / 2
    slide.shapes.add_picture(str(path), int(x), int(y), int(draw_w), int(draw_h))


def _draw_title(
    slide,
    title: str,
    theme: GridTheme,
    *,
    x: int = 0,
    y: int = 0,
    width: int = SLIDE_W,
) -> None:
    if not title.strip():
        return
    size = theme.title_pt
    if len(title) > 22:
        size -= 2
    if len(title) > 32:
        size -= 2
    frame = _frame(
        slide,
        x,
        y,
        width,
        EMU_PER_CELL * TITLE_ROWS,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    paragraph = frame.paragraphs[0]
    paragraph.line_spacing = 1.2
    run = paragraph.add_run()
    run.text = title
    _style(run, size, theme, bold=True, color=theme.title)


def _flow_rows(block: GridBlock, width: int, theme: GridTheme) -> int:
    need = _flow_h(block.flow, width, theme.body_pt, theme.line_spacing)
    return max(1, math.ceil(need / EMU_PER_CELL))


def _image_rows(block: GridBlock, width: int, theme: GridTheme, root: Path) -> int:
    path = (root / block.path).resolve() if block.path else None
    aspect = _aspect(path) if path and path.is_file() else 4 / 3
    caption = theme.caption_pt * 1.8 * EMU_PER_PT if block.caption else 0
    need = (width - 2 * PAD) / aspect + caption + PAD
    return max(1, math.ceil(need / EMU_PER_CELL))


def _table_rows(block: GridBlock) -> int:
    need = (1 + len(block.rows)) * int(0.36 * 914400) + 2 * PAD
    return max(2, math.ceil(need / EMU_PER_CELL))


def _block_rows(block: GridBlock, width: int, theme: GridTheme, root: Path) -> int:
    if block.kind == IMAGE:
        return _image_rows(block, width, theme, root)
    if block.kind == TABLE:
        return _table_rows(block)
    return _flow_rows(block, width, theme)


def _auto_placement(
    slide: GridSlide, blocks: list[GridBlock], theme: GridTheme, root: Path
) -> list[tuple[GridBlock, Box]]:
    """未定位的块按内容形态自动组合版面。"""
    if not blocks:
        return []
    top = TITLE_ROWS + 1 if slide.title else 1
    rows = GRID_ROWS - top + 1
    full_w = GRID_COLS * EMU_PER_CELL

    texts = [b for b in blocks if b.kind in ("text", "quote", "code", FORMULA)]
    images = [b for b in blocks if b.kind == IMAGE]
    tables = [b for b in blocks if b.kind == TABLE]

    # 左文右图：一个文本块 + 一张图
    if len(texts) == 1 and len(images) == 1 and not tables:
        img_rows = min(rows, _image_rows(images[0], 7 * EMU_PER_CELL, theme, root))
        return [
            (texts[0], Box(1, top, 8, rows)),
            (images[0], Box(10, top + (rows - img_rows) // 2, 7, img_rows)),
        ]

    # 纯图片页
    if images and not texts and not tables:
        if len(images) == 1:
            return [(images[0], Box(2, top, 14, rows))]
        if len(images) == 2:
            return [(images[0], Box(1, top, 7, rows)), (images[1], Box(9, top, 8, rows))]
        half = max(1, rows // 2)
        out: list[tuple[GridBlock, Box]] = []
        cols = 2
        per_row = max(1, math.ceil(len(images) / cols))
        cell_rows = max(2, rows // per_row)
        for index, image in enumerate(images):
            row, col = divmod(index, cols)
            x = 1 + col * 8
            out.append((image, Box(x, min(top + row * cell_rows, GRID_ROWS), 8, cell_rows)))
        return out

    # 纯表格页
    if tables and not texts and not images:
        if len(tables) == 1:
            return [(tables[0], Box(1, top, GRID_COLS, rows))]
        cursor = top
        out = []
        for table in tables:
            span = min(_table_rows(table), GRID_ROWS - cursor + 1)
            out.append((table, Box(1, cursor, GRID_COLS, span)))
            cursor = min(GRID_ROWS, cursor + span)
        return out

    # 文本 + 图片/表格：文本在上，其余在下方纵向堆叠
    cursor = top
    out = []
    if texts:
        text_rows = min(
            max(1, rows - 2),
            sum(_flow_rows(block, full_w, theme) for block in texts),
        )
        out.append((texts[0], Box(1, cursor, GRID_COLS, text_rows)))
        cursor += text_rows
        for extra in texts[1:]:
            span = min(_flow_rows(extra, full_w, theme), GRID_ROWS - cursor + 1)
            out.append((extra, Box(1, cursor, GRID_COLS, span)))
            cursor = min(GRID_ROWS, cursor + span)

    for block in [*tables, *images]:
        span = _block_rows(block, full_w, theme, root)
        span = min(span, GRID_ROWS - cursor + 1)
        out.append((block, Box(1, cursor, GRID_COLS, span)))
        cursor = min(GRID_ROWS, cursor + span)
    return out


def _widen_lone_text(slide: GridSlide) -> None:
    """整页只有一个正文文字块时铺满内容区，左右各留 1 列。"""
    content = [
        block
        for block in slide.blocks
        if not (block.box and block.box.row == 1 and block.box.col >= 9)
    ]
    if len(content) != 1:
        return
    block = content[0]
    box = block.box
    if block.kind != TEXT or box is None:
        return
    width = box.colspan or (GRID_COLS - box.col + 1)
    if box.col != 2 or width != GRID_COLS - 2:
        block.box = replace(box, col=2, colspan=GRID_COLS - 2)


def _render_slide(
    pptx_slide,
    slide: GridSlide,
    theme: GridTheme,
    root: Path,
    report: RenderReport,
    logo: Path | None = None,
    header: str | None = None,
) -> None:
    _widen_lone_text(slide)
    fill = pptx_slide.background.fill
    fill.solid()
    fill.fore_color.rgb = theme.background

    if slide.title.strip().lower() not in COVER_TITLES:
        title_y = LOGO_ROWS * EMU_PER_CELL if logo else 0
        _draw_title(pptx_slide, slide.title, theme, y=title_y)

    positioned = [block for block in slide.blocks if block.box is not None]
    floating = [block for block in slide.blocks if block.box is None]
    placements = _auto_placement(slide, floating, theme, root) + [
        (block, block.box) for block in positioned if block.box is not None
    ]
    limit = SLIDE_H - BOTTOM_SAFE
    placements = [
        (block, box.with_bottom_limit(limit)) for block, box in placements
    ]

    tag = f"P{slide.number:02d}"
    for block, box in placements:
        report.blocks += 1
        if not block.is_positioned:
            report.auto_blocks += 1
        if block.kind == IMAGE:
            _draw_image(pptx_slide, block, box, theme, root, report, tag)
        elif block.kind == TABLE:
            _draw_table(pptx_slide, block, box, theme, report, tag)
        elif block.kind == CODE:
            _draw_code(pptx_slide, block, box, theme, report, tag)
        elif block.kind == FORMULA:
            _draw_formula(pptx_slide, block, box, theme, report, tag)
        elif block.kind == QUOTE:
            _draw_quote(pptx_slide, block, box, theme, report, tag)
        else:
            _draw_flow(pptx_slide, block, box, theme, report, tag)

    if logo is not None:
        _draw_header_rule(pptx_slide, theme)
    if header:
        _draw_header_text(pptx_slide, header, theme)
    if logo is not None:
        _draw_logo(pptx_slide, logo)

    if slide.notes:
        pptx_slide.notes_slide.notes_text_frame.text = slide.notes


# --------------------------------------------------------------------------
# 入口
# --------------------------------------------------------------------------
def render_deck(
    deck: GridDeck,
    out_path: str | Path,
    *,
    theme: GridTheme = DEFAULT_THEME,
    base_dir: str | Path | None = None,
    logo: str | Path | bool | None = None,
    header: str | None = None,
) -> tuple[Path, RenderReport]:
    root = Path(base_dir) if base_dir else (deck.source.parent if deck.source else Path.cwd())
    if logo is False:
        logo_path = None
    elif logo is None:
        logo_path = DEFAULT_LOGO if DEFAULT_LOGO.is_file() else None
    else:
        logo_path = Path(logo).resolve()
        if not logo_path.is_file():
            raise ValueError(f"找不到 logo 文件：{logo}")
    presentation = Presentation()
    presentation.slide_width = Emu(SLIDE_W)
    presentation.slide_height = Emu(SLIDE_H)
    blank = presentation.slide_layouts[6]
    report = RenderReport()

    for slide in deck.slides:
        pptx_slide = presentation.slides.add_slide(blank)
        _render_slide(pptx_slide, slide, theme, root, report, logo_path, header)
        report.slides += 1

    presentation.core_properties.title = deck.title
    target = Path(out_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(target))
    return target, report


def render_file(
    md_path: str | Path,
    out_path: str | Path | None = None,
    *,
    theme: GridTheme = DEFAULT_THEME,
    logo: str | Path | bool | None = None,
    header: str | None = None,
) -> tuple[Path, RenderReport]:
    from .grid import load_grid

    source = Path(md_path)
    deck = load_grid(source)
    target = Path(out_path) if out_path else source.with_suffix(".pptx")
    return render_deck(
        deck,
        target,
        theme=theme,
        base_dir=source.parent,
        logo=logo,
        header=header,
    )


def make_theme(accent: str | None = None) -> GridTheme:
    if not accent:
        return DEFAULT_THEME
    value = accent.strip().lstrip("#").upper()
    if len(value) != 6:
        raise ValueError(f"--accent 需要 6 位十六进制颜色，收到 {accent!r}")
    rgb = RGBColor.from_string(value)
    return replace(DEFAULT_THEME, accent=rgb, table_head=rgb)
