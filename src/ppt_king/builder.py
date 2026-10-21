"""版式规格 -> PPTX 的确定性渲染器（不调用模型）。

所有几何计算基于 16:9（13.333 x 7.5 英寸）= 12192000 x 6858000 EMU，
与常见宽屏演示文稿一致。文字高度用字符宽度估算，必要时按比例缩小字号。
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from .spec import Block, DeckSpec, SlideSpec

SLIDE_W = 12192000
SLIDE_H = 6858000
MARGIN_X = Inches(0.62)
MARGIN_BOTTOM = Inches(0.55)

EMU_PER_PT = 12700


@dataclass(frozen=True)
class Theme:
    font: str = "Microsoft YaHei"
    text: RGBColor = RGBColor(0x1F, 0x29, 0x37)
    muted: RGBColor = RGBColor(0x6B, 0x72, 0x80)
    accent: RGBColor = RGBColor(0x1D, 0x4E, 0xD8)
    accent_soft: RGBColor = RGBColor(0xE8, 0xEF, 0xFB)
    surface: RGBColor = RGBColor(0xF8, 0xFA, 0xFC)
    border: RGBColor = RGBColor(0xD8, 0xDE, 0xE8)
    background: RGBColor = RGBColor(0xFF, 0xFF, 0xFF)


DEFAULT_THEME = Theme()

BASE_TEXT_PT = 16


# --------------------------------------------------------------------------
# 基础工具
# --------------------------------------------------------------------------
def _style_run(run, size_pt: float, theme: Theme, *, bold=False, color=None, italic=False) -> None:
    font = run.font
    font.size = Pt(size_pt)
    font.bold = bold
    font.italic = italic
    font.color.rgb = color or theme.text
    font.name = theme.font
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rpr, qn("a:ea"))
    ea.set("typeface", theme.font)


def _textbox(
    slide,
    left: int,
    top: int,
    width: int,
    height: int,
    *,
    anchor=MSO_ANCHOR.TOP,
):
    shape = slide.shapes.add_textbox(int(left), int(top), int(width), int(height))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    return shape, frame


def _rect(slide, left: int, top: int, width: int, height: int, color: RGBColor):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, int(left), int(top), int(width), int(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    if shape.has_text_frame:
        shape.text_frame.word_wrap = False
    return shape


def _bullet_marker(paragraph, char: str, level: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    ppr.set("marL", str(int(Inches(0.26 + 0.30 * level))))
    ppr.set("indent", str(int(-Inches(0.24))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
        for node in ppr.findall(qn(tag)):
            ppr.remove(node)
    font_node = etree.SubElement(ppr, qn("a:buFont"))
    font_node.set("typeface", "Arial")
    char_node = etree.SubElement(ppr, qn("a:buChar"))
    char_node.set("char", char)


BULLET_CHARS = ("•", "–", "·", "◦")


def _text_units(text: str) -> float:
    total = 0.0
    for char in text:
        code = ord(char)
        if char in " \t":
            total += 0.32
        elif code < 0x2E80:
            total += 0.55
        else:
            total += 1.0
    return total or 1.0


def _line_count(text: str, width_emu: int, size_pt: float) -> int:
    width_pt = max(1.0, (width_emu - Inches(0.08)) / EMU_PER_PT)
    per_line = max(1.0, width_pt / max(1.0, size_pt))
    return max(1, math.ceil(_text_units(text) / per_line))


def _text_height(text: str, width_emu: int, size_pt: float, factor: float = 1.4) -> int:
    return int(_line_count(text, width_emu, size_pt) * size_pt * factor * EMU_PER_PT)


def _aspect(path: Path) -> float:
    try:
        with Image.open(path) as image:
            width, height = image.size
        if width and height:
            return width / height
    except Exception:  # pragma: no cover - 图片损坏时退化为 4:3
        pass
    return 4 / 3


# --------------------------------------------------------------------------
# 块高度估算
# --------------------------------------------------------------------------
def _block_height(block: Block, width: int, size_pt: float, theme: Theme) -> int:
    if block.type == "heading":
        return _text_height(block.text, width, size_pt + 2, 1.5) + int(Inches(0.10))
    if block.type == "paragraph":
        return _text_height(block.text, width, size_pt) + int(Inches(0.04))
    if block.type == "bullets":
        total = 0
        for item in block.items:
            inner = width - int(Inches(0.32 * (item.level + 1)))
            total += _text_height(item.text, inner, size_pt - 1.5 * item.level, 1.45)
        return total + int(Inches(0.04))
    if block.type == "quote":
        return _text_height(block.text, width - Inches(0.35), size_pt - 1, 1.45) + int(
            Inches(0.16)
        )
    if block.type == "table":
        rows = 1 + len(block.rows)
        return int(Inches(0.42) + (rows - 1) * Inches(0.34)) + int(Inches(0.12))
    return int(Inches(0.3))


# --------------------------------------------------------------------------
# 块渲染
# --------------------------------------------------------------------------
def _draw_text_block(
    slide, block: Block, left: int, top: int, width: int, height: int, size_pt: float, theme: Theme
) -> None:
    if block.type == "heading":
        _, frame = _textbox(slide, left, top, width, height)
        paragraph = frame.paragraphs[0]
        paragraph.line_spacing = 1.25
        run = paragraph.add_run()
        run.text = block.text
        _style_run(run, size_pt + 2, theme, bold=True, color=theme.accent)
        return

    if block.type == "paragraph":
        _, frame = _textbox(slide, left, top, width, height)
        paragraph = frame.paragraphs[0]
        paragraph.line_spacing = 1.4
        run = paragraph.add_run()
        run.text = block.text
        _style_run(run, size_pt, theme)
        return

    if block.type == "bullets":
        _, frame = _textbox(slide, left, top, width, height)
        for index, item in enumerate(block.items):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.line_spacing = 1.35
            paragraph.space_before = Pt(2 if index else 0)
            run = paragraph.add_run()
            run.text = item.text
            _style_run(run, size_pt - 1.5 * item.level, theme)
            _bullet_marker(paragraph, BULLET_CHARS[min(item.level, len(BULLET_CHARS) - 1)], item.level)
        return

    if block.type == "quote":
        bar_width = Inches(0.06)
        _rect(slide, left, top + Inches(0.04), bar_width, height - Inches(0.10), theme.accent)
        _, frame = _textbox(slide, left + Inches(0.18), top, width - Inches(0.22), height)
        paragraph = frame.paragraphs[0]
        paragraph.line_spacing = 1.4
        run = paragraph.add_run()
        run.text = block.text
        _style_run(run, size_pt - 1, theme, color=theme.muted)
        return

    if block.type == "table":
        _draw_table(slide, block, left, top, width, height, theme)


def _column_weights(header: list[str], rows: list[list[str]]) -> list[float]:
    weights: list[float] = []
    for index, cell in enumerate(header):
        longest = _text_units(cell)
        for row in rows:
            if index < len(row):
                longest = max(longest, _text_units(row[index]))
        weights.append(max(2.2, min(longest, 18.0)))
    total = sum(weights) or 1.0
    return [weight / total for weight in weights]


def _draw_table(
    slide, block: Block, left: int, top: int, width: int, height: int, theme: Theme
) -> None:
    rows = [block.header, *block.rows]
    cols = len(block.header)
    row_height = min(Inches(0.42), max(Inches(0.26), height // max(1, len(rows))))
    table_height = row_height * len(rows)

    shape = slide.shapes.add_table(
        len(rows), cols, int(left), int(top), int(width), int(table_height)
    )
    table = shape.table
    table.first_row = True
    table.horz_banding = False

    weights = _column_weights(block.header, block.rows)
    for index, weight in enumerate(weights):
        table.columns[index].width = Emu(int(width * weight))
    for row in table.rows:
        row.height = Emu(int(row_height))

    size_pt = 12 if cols <= 4 else 10
    if len(rows) > 8:
        size_pt -= 1
    if row_height < Inches(0.30):
        size_pt -= 1

    for row_index, row in enumerate(rows):
        for col_index in range(cols):
            cell = table.cell(row_index, col_index)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if row_index == 0:
                cell.fill.fore_color.rgb = theme.accent_soft
            else:
                cell.fill.fore_color.rgb = (
                    theme.surface if row_index % 2 == 0 else theme.background
                )
            text = row[col_index] if col_index < len(row) else ""
            frame = cell.text_frame
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.CENTER if col_index else PP_ALIGN.LEFT
            run = paragraph.add_run()
            run.text = text
            _style_run(run, max(8.0, size_pt), theme, bold=row_index == 0)


def _draw_image(
    slide,
    path: Path,
    left: int,
    top: int,
    width: int,
    height: int,
    caption: str,
    theme: Theme,
    *,
    caption_pt: float = 10.5,
) -> int:
    """在矩形内等比放置图片，返回实际占用的总高度。"""
    caption_h = _text_height(caption, width, caption_pt, 1.4) if caption else 0
    box_h = max(int(Inches(0.5)), height - caption_h - int(Inches(0.04)))
    aspect = _aspect(path)
    draw_w = box_h * aspect
    draw_h = box_h
    if draw_w > width:
        draw_w = width
        draw_h = width / aspect
    x = left + (width - draw_w) / 2
    y = top + (box_h - draw_h) / 2
    try:
        slide.shapes.add_picture(str(path), int(x), int(y), int(draw_w), int(draw_h))
    except Exception:  # pragma: no cover - 图片无法解码
        return caption_h
    if caption:
        _, frame = _textbox(slide, left, int(top + box_h), width, caption_h)
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = caption
        _style_run(run, caption_pt, theme, color=theme.muted)
    return box_h + caption_h


# --------------------------------------------------------------------------
# 页级版式
# --------------------------------------------------------------------------
def _paint_background(slide, theme: Theme) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = theme.background


def _title_band(slide, title: str, theme: Theme) -> int:
    if not title.strip():
        return Inches(0.45)
    size = 30.0
    if len(title) > 26:
        size = 25.0
    if len(title) > 36:
        size = 21.0
    height = Inches(0.85)
    _, frame = _textbox(slide, MARGIN_X, Inches(0.36), SLIDE_W - 2 * MARGIN_X, height)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.line_spacing = 1.15
    run = paragraph.add_run()
    run.text = title or " "
    _style_run(run, size, theme, bold=True)
    _rect(slide, MARGIN_X, Inches(1.23), Inches(0.92), Inches(0.055), theme.accent)
    return Inches(1.58)


def _body_rect(top: int) -> tuple[int, int, int, int]:
    return (
        int(MARGIN_X),
        int(top),
        int(SLIDE_W - 2 * MARGIN_X),
        int(SLIDE_H - MARGIN_BOTTOM - top),
    )


def _flow(slide, blocks: list[Block], rect, theme: Theme, assets: dict[str, Path]) -> None:
    left, top, width, height = rect
    texts = [block for block in blocks if block.type != "image"]
    images = [block for block in blocks if block.type == "image"]

    gap = int(Inches(0.14))
    size_pt = float(BASE_TEXT_PT)
    heights: list[int] = []
    for scale in (1.0, 0.92, 0.84, 0.76, 0.68):
        size_pt = BASE_TEXT_PT * scale
        heights = [_block_height(block, width, size_pt, theme) for block in texts]
        used = sum(heights) + gap * max(0, len(texts) - 1)
        if not images or height - used >= int(Inches(1.4)) * len(images):
            break

    used = sum(heights) + gap * max(0, len(texts) - 1)
    image_space = max(0, height - used - (gap if texts and images else 0))
    per_image = int(image_space / len(images)) if images else 0

    cursor = top
    for block, block_h in zip(texts, heights):
        _draw_text_block(slide, block, left, cursor, width, block_h, size_pt, theme)
        cursor += block_h + gap

    if images and per_image > 0:
        cursor += gap if texts else 0
        for block in images:
            path = assets.get(block.ref)
            if path is None or not path.is_file():
                cursor += gap
                continue
            _draw_image(slide, path, left, cursor, width, per_image, block.caption, theme)
            cursor += per_image + gap


def _two_column(
    slide,
    blocks: list[Block],
    rect,
    theme: Theme,
    assets: dict[str, Path],
    *,
    images_left: bool,
) -> None:
    left, top, width, height = rect
    gap = int(Inches(0.34))
    text_width = int((width - gap) * 0.52)
    image_width = width - gap - text_width
    texts = [block for block in blocks if block.type != "image"]
    images = [block for block in blocks if block.type == "image"]

    if images_left:
        image_left, text_left = left, left + image_width + gap
    else:
        text_left, image_left = left, left + text_width + gap

    _flow(slide, texts, (text_left, top, text_width, height), theme, assets)

    if not images:
        return
    per_image = max(int(Inches(1.2)), height // len(images) - gap)
    cursor = top
    for block in images:
        path = assets.get(block.ref)
        if path is None or not path.is_file():
            cursor += per_image + gap
            continue
        _draw_image(slide, path, image_left, cursor, image_width, per_image, block.caption, theme)
        cursor += per_image + gap


def _gallery(slide, blocks: list[Block], rect, theme: Theme, assets: dict[str, Path]) -> None:
    left, top, width, height = rect
    images = [block for block in blocks if block.type == "image"]
    others = [block for block in blocks if block.type != "image"]
    gap = int(Inches(0.20))

    text_height = 0
    if others:
        text_height = min(int(height * 0.34), sum(
            _block_height(block, width, BASE_TEXT_PT * 0.85, theme) for block in others
        ) + gap)
        _flow(slide, others, (left, top, width, text_height), theme, assets)
        top += text_height + gap
        height -= text_height + gap

    cols = 1 if len(images) == 1 else 2
    rows = max(1, math.ceil(len(images) / cols))
    cell_w = int((width - gap * (cols - 1)) / cols)
    cell_h = int((height - gap * (rows - 1)) / rows)

    for index, block in enumerate(images):
        row, col = divmod(index, cols)
        x = left + col * (cell_w + gap)
        y = top + row * (cell_h + gap)
        path = assets.get(block.ref)
        if path is None or not path.is_file():
            continue
        _draw_image(slide, path, x, y, cell_w, cell_h, block.caption, theme, caption_pt=10)


def _layout_cover(slide, spec: SlideSpec, theme: Theme, assets: dict[str, Path]) -> None:
    texts: list[str] = []
    if spec.subtitle:
        texts.append(spec.subtitle)
    for block in spec.blocks:
        if block.type in ("paragraph", "heading"):
            texts.append(block.text)
        elif block.type == "bullets":
            texts.extend(item.text for item in block.items)
        elif block.type == "quote":
            texts.append(block.text)
    images = [block for block in spec.blocks if block.type == "image"]

    has_text = bool(texts)
    title_top = Inches(2.05) if has_text or images else Inches(2.6)
    _rect(
        slide,
        int(SLIDE_W / 2 - Inches(0.5)),
        int(title_top - Inches(0.30)),
        int(Inches(1.0)),
        int(Inches(0.06)),
        theme.accent,
    )

    title_size = 42.0
    if len(spec.title) > 18:
        title_size = 34.0
    if len(spec.title) > 30:
        title_size = 28.0
    _, frame = _textbox(
        slide, MARGIN_X, title_top, SLIDE_W - 2 * MARGIN_X, Inches(1.5)
    )
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.line_spacing = 1.2
    run = paragraph.add_run()
    run.text = spec.title
    _style_run(run, title_size, theme, bold=True)

    cursor = title_top + Inches(1.35)
    if texts:
        body_height = min(int(Inches(2.0)), Inches(0.4) * len(texts) + Inches(0.3))
        _, frame = _textbox(
            slide, Inches(1.2), cursor, SLIDE_W - 2 * Inches(1.2), body_height
        )
        for index, text in enumerate(texts):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.line_spacing = 1.35
            paragraph.space_before = Pt(6 if index else 0)
            run = paragraph.add_run()
            run.text = text
            _style_run(run, 18 if index else 20, theme, color=theme.text if index == 0 else theme.muted)
        cursor += body_height

    if images:
        strip_top = max(int(cursor + Inches(0.25)), int(SLIDE_H - Inches(1.85)))
        strip_h = int(SLIDE_H - MARGIN_BOTTOM - strip_top)
        cell_w = int((SLIDE_W - 2 * MARGIN_X) / len(images))
        for index, block in enumerate(images):
            path = assets.get(block.ref)
            if path is None or not path.is_file():
                continue
            _draw_image(
                slide,
                path,
                int(MARGIN_X + index * cell_w),
                strip_top,
                cell_w - int(Inches(0.15)),
                strip_h,
                "",
                theme,
            )


def _layout_section(slide, spec: SlideSpec, theme: Theme) -> None:
    _rect(
        slide,
        int(SLIDE_W / 2 - Inches(0.5)),
        int(Inches(2.55)),
        int(Inches(1.0)),
        int(Inches(0.06)),
        theme.accent,
    )
    _, frame = _textbox(
        slide, MARGIN_X, Inches(2.75), SLIDE_W - 2 * MARGIN_X, Inches(1.4)
    )
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = spec.title
    _style_run(run, 34, theme, bold=True)
    if spec.subtitle:
        _, frame = _textbox(
            slide, Inches(1.2), Inches(4.15), SLIDE_W - 2 * Inches(1.2), Inches(0.6)
        )
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = spec.subtitle
        _style_run(run, 16, theme, color=theme.muted)


def _slide_number(slide, number: int, theme: Theme, layout: str) -> None:
    if not number or layout in ("title", "section"):
        return
    _, frame = _textbox(
        slide,
        SLIDE_W - MARGIN_X - Inches(0.8),
        SLIDE_H - MARGIN_BOTTOM,
        Inches(0.8),
        Inches(0.3),
    )
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.RIGHT
    run = paragraph.add_run()
    run.text = str(number)
    _style_run(run, 10, theme, color=theme.muted)


def _render_slide(slide, spec: SlideSpec, theme: Theme, assets: dict[str, Path]) -> None:
    _paint_background(slide, theme)

    if spec.layout == "title":
        _layout_cover(slide, spec, theme, assets)
    elif spec.layout == "section":
        _layout_section(slide, spec, theme)
    else:
        top = _title_band(slide, spec.title, theme)
        rect = _body_rect(top)
        if spec.layout == "gallery":
            _gallery(slide, spec.blocks, rect, theme, assets)
        elif spec.layout == "two_column":
            _two_column(slide, spec.blocks, rect, theme, assets, images_left=False)
        elif spec.layout == "image_left":
            _two_column(slide, spec.blocks, rect, theme, assets, images_left=True)
        else:
            _flow(slide, spec.blocks, rect, theme, assets)

    if spec.notes:
        slide.notes_slide.notes_text_frame.text = spec.notes
    _slide_number(slide, spec.number, theme, spec.layout)


# --------------------------------------------------------------------------
# 对外入口
# --------------------------------------------------------------------------
def resolve_assets(spec: DeckSpec, base_dir: Path) -> dict[int, dict[str, Path]]:
    resolved: dict[int, dict[str, Path]] = {}
    for slide in spec.slides:
        resolved[slide.number] = {
            ref: (base_dir / rel) for ref, rel in slide.assets.items()
        }
    return resolved


def build_deck(
    spec: DeckSpec,
    out_path: str | Path,
    *,
    theme: Theme = DEFAULT_THEME,
    base_dir: str | Path | None = None,
) -> Path:
    """把 DeckSpec 渲染成 .pptx。``base_dir`` 用于解析 spec 里的相对资源路径。"""
    root = Path(base_dir) if base_dir else Path.cwd()
    assets = resolve_assets(spec, root)

    presentation = Presentation()
    presentation.slide_width = Emu(SLIDE_W)
    presentation.slide_height = Emu(SLIDE_H)
    blank = presentation.slide_layouts[6]

    for slide_spec in spec.slides:
        slide = presentation.slides.add_slide(blank)
        _render_slide(slide, slide_spec, theme, assets.get(slide_spec.number, {}))

    presentation.core_properties.title = spec.title
    target = Path(out_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(target))
    return target
