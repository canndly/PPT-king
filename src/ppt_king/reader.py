"""PPTX -> IR 的确定性抽取（不含 AI）。"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from .assets import AssetSink
from .ir import (
    ChartBlock,
    Deck,
    EmbedBlock,
    ImageBlock,
    Paragraph,
    Run,
    ShapeBase,
    Slide,
    TableBlock,
    TextBlock,
)

ROW_TOLERANCE = 45720  # 0.05 英寸，用于把同一水平行的形状归组
MONO_FONTS = {
    "consolas",
    "courier new",
    "courier",
    "monaco",
    "menlo",
    "dejavu sans mono",
    "source code pro",
    "jetbrains mono",
}
SKIP_PLACEHOLDERS = {"SLIDE_NUMBER", "FOOTER", "DATE", "HEADER"}
TITLE_PLACEHOLDERS = {"TITLE", "CENTER_TITLE", "VERTICAL_TITLE"}
SUBTITLE_PLACEHOLDERS = {"SUBTITLE"}

CHART_TYPE_NAMES = {
    "COLUMN_CLUSTERED": "柱状图",
    "COLUMN_STACKED": "堆积柱状图",
    "COLUMN_STACKED_100": "百分比堆积柱状图",
    "BAR_CLUSTERED": "条形图",
    "BAR_STACKED": "堆积条形图",
    "LINE": "折线图",
    "LINE_MARKERS": "带标记折线图",
    "PIE": "饼图",
    "DOUGHNUT": "环形图",
    "AREA": "面积图",
    "XY_SCATTER": "散点图",
    "XY_SCATTER_LINES": "带线散点图",
    "RADAR": "雷达图",
    "BUBBLE": "气泡图",
    "STOCK_HLC": "股价图",
}


def _enum_name(value: Any) -> str:
    name = getattr(value, "name", None)
    if isinstance(name, str):
        return name
    text = str(value)
    return text.split(" (", 1)[0].strip()


def _on(value: Any) -> bool:
    return value in (True, 1, "1", "true")


def _sort_key(shape: Any) -> tuple[int, int]:
    return (shape.top or 0, shape.left or 0)


def reading_order(shapes: Iterable[Any]) -> list[Any]:
    """按坐标还原阅读顺序：先按行（top 容差），行内按 left。"""
    ordered = sorted(shapes, key=_sort_key)
    rows: list[tuple[int, list[Any]]] = []
    for shape in ordered:
        top = shape.top or 0
        if rows and abs(top - rows[-1][0]) < ROW_TOLERANCE:
            rows[-1][1].append(shape)
        else:
            rows.append((top, [shape]))
    result: list[Any] = []
    for _, row in rows:
        result.extend(sorted(row, key=lambda s: s.left or 0))
    return result


def _paragraph_runs(paragraph: Any) -> list[Run]:
    runs: list[Run] = []
    for child in paragraph._p.iterchildren():
        tag = child.tag
        if tag == qn("a:r"):
            node = child.find(qn("a:t"))
            text = (node.text or "") if node is not None else ""
            if not text:
                continue
            bold = italic = code = False
            rpr = child.find(qn("a:rPr"))
            if rpr is not None:
                bold = _on(rpr.get("b"))
                italic = _on(rpr.get("i"))
                latin = rpr.find(qn("a:latin"))
                if latin is not None and (latin.get("typeface") or "").lower() in MONO_FONTS:
                    code = True
            runs.append(Run(text.replace("\x0b", "\n"), bold, italic, code))
        elif tag == qn("a:br"):
            runs.append(Run("\n"))
        elif tag == qn("a:fld"):
            node = child.find(qn("a:t"))
            text = (node.text or "") if node is not None else ""
            if text:
                runs.append(Run(text.replace("\x0b", "\n")))
    return runs


def _bullet(paragraph: Any) -> str | None:
    ppr = paragraph._p.find(qn("a:pPr"))
    if ppr is None:
        return None
    if ppr.find(qn("a:buNone")) is not None:
        return None
    if ppr.find(qn("a:buAutoNum")) is not None:
        return "number"
    if ppr.find(qn("a:buChar")) is not None:
        return "char"
    return None


def _paragraphs(text_frame: Any) -> list[Paragraph]:
    result: list[Paragraph] = []
    for para in text_frame.paragraphs:
        runs = _paragraph_runs(para)
        paragraph = Paragraph(runs=runs, level=int(para.level or 0), bullet=_bullet(para))
        if paragraph.is_empty and result:
            continue
        result.append(paragraph)
    while result and result[-1].is_empty:
        result.pop()
    return result


def _cell_text(cell: Any) -> str:
    parts: list[str] = []
    for para in cell.text_frame.paragraphs:
        text = "".join(run.text for run in _paragraph_runs(para))
        parts.append(text.strip())
    joined = "\n".join(part for part in parts)
    return joined.strip("\n")


def _table_block(shape: Any) -> TableBlock:
    table = shape.table
    rows: list[list[str]] = []
    merges: list[tuple[int, int, int, int]] = []
    spanned: set[tuple[int, int]] = set()
    for row_index, row in enumerate(table.rows):
        cells: list[str] = []
        for col_index, cell in enumerate(row.cells):
            if cell.is_spanned:
                spanned.add((row_index, col_index))
                cells.append("")
            else:
                cells.append(_cell_text(cell))
            if cell.is_merge_origin and (cell.span_width > 1 or cell.span_height > 1):
                merges.append((row_index, col_index, cell.span_height, cell.span_width))
        rows.append(cells)
    return TableBlock(
        name=shape.name,
        top=shape.top or 0,
        left=shape.left or 0,
        width=shape.width or 0,
        height=shape.height or 0,
        rows=rows,
        merges=merges,
        spanned=spanned,
    )


def _fmt_number(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.6g}"
    return str(value)


def _chart_block(shape: Any) -> ChartBlock:
    chart = shape.chart
    raw_type = _enum_name(chart.chart_type)
    chart_type = CHART_TYPE_NAMES.get(raw_type, raw_type)

    title = ""
    try:
        if chart.has_title:
            title = chart.chart_title.text_frame.text.strip()
    except Exception:  # pragma: no cover - 个别畸形图表
        title = ""

    categories: list[str] = []
    series: list[tuple[str, list[str]]] = []
    try:
        plots = list(chart.plots)
    except Exception:  # pragma: no cover
        plots = []

    try:
        first = plots[0] if plots else None
        if first is not None and first.categories is not None:
            categories = [_fmt_number(item) for item in first.categories]
    except Exception:  # pragma: no cover - 散点图等无分类
        categories = []

    for plot in plots:
        for item in plot.series:
            try:
                values = [_fmt_number(value) for value in item.values]
            except Exception:  # pragma: no cover
                values = []
            name = ""
            try:
                name = str(item.name or "").strip()
            except Exception:  # pragma: no cover
                name = ""
            series.append((name or f"系列{len(series) + 1}", values))

    return ChartBlock(
        name=shape.name,
        top=shape.top or 0,
        left=shape.left or 0,
        width=shape.width or 0,
        height=shape.height or 0,
        chart_type=chart_type,
        title=title,
        categories=categories,
        series=series,
    )


def _media_relation(shape: Any) -> tuple[str | None, str]:
    element = shape._element
    for tag, kind in ((qn("a:videoFile"), "video"), (qn("a:audioFile"), "audio")):
        node = element.find(".//" + tag)
        if node is not None:
            rel_id = node.get(qn("r:link")) or node.get(qn("r:embed"))
            if rel_id:
                return rel_id, kind
    node = element.find(".//" + qn("p:oleObj"))
    if node is not None:
        rel_id = node.get(qn("r:id"))
        if rel_id:
            return rel_id, "ole"
    return None, "file"


def _embed_block(shape: Any, slide_no: int, sink: AssetSink) -> EmbedBlock | None:
    rel_id, kind = _media_relation(shape)
    if not rel_id:
        return None
    try:
        part = shape.part.related_part(rel_id)
        blob = part.blob
    except Exception:  # pragma: no cover - 外部链接对象没有内嵌数据
        return None
    ext = Path(str(part.partname)).suffix.lstrip(".")
    asset_id, rel_path = sink.add_file(blob, ext, slide_no, default_ext="bin")
    label = (shape.name or "").strip() or f"{kind} 对象"
    return EmbedBlock(
        name=shape.name,
        top=shape.top or 0,
        left=shape.left or 0,
        width=shape.width or 0,
        height=shape.height or 0,
        asset_id=asset_id,
        rel_path=rel_path,
        kind=kind,
        label=label,
    )


def _image_alt(shape: Any) -> str:
    try:
        cnv = shape._element.find(".//" + qn("p:cNvPr"))
        if cnv is not None:
            descr = (cnv.get("descr") or "").strip()
            if descr:
                return descr
    except Exception:  # pragma: no cover
        pass
    return (shape.name or "").strip()


def _geometry(shape: Any) -> dict[str, int]:
    return {
        "name": shape.name,
        "top": shape.top or 0,
        "left": shape.left or 0,
        "width": shape.width or 0,
        "height": shape.height or 0,
    }


def _placeholder_name(shape: Any) -> str:
    try:
        return _enum_name(shape.placeholder_format.type)
    except Exception:  # pragma: no cover
        return ""


def _shape_blocks(shape: Any, slide_no: int, sink: AssetSink) -> list[ShapeBase]:
    if shape.is_placeholder:
        ph = _placeholder_name(shape)
        if ph in SKIP_PLACEHOLDERS:
            return []

    shape_type = shape.shape_type

    if shape_type == MSO_SHAPE_TYPE.GROUP:
        blocks: list[ShapeBase] = []
        for child in reading_order(shape.shapes):
            blocks.extend(_shape_blocks(child, slide_no, sink))
        return blocks

    if shape.has_table:
        return [_table_block(shape)]

    if getattr(shape, "has_chart", False):
        return [_chart_block(shape)]

    if shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            image = shape.image
        except Exception:  # pragma: no cover - 链接图片无内嵌数据
            return []
        asset_id, rel_path = sink.add_image(image.blob, image.ext, slide_no)
        return [
            ImageBlock(
                **_geometry(shape),
                asset_id=asset_id,
                rel_path=rel_path,
                alt=_image_alt(shape),
            )
        ]

    if shape_type in (MSO_SHAPE_TYPE.MEDIA, MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT):
        embed = _embed_block(shape, slide_no, sink)
        return [embed] if embed else []

    if shape.has_text_frame:
        paragraphs = _paragraphs(shape.text_frame)
        if not paragraphs or not any(not p.is_empty for p in paragraphs):
            return []
        ph = _placeholder_name(shape) if shape.is_placeholder else ""
        if ph in TITLE_PLACEHOLDERS:
            role = "title"
        elif ph in SUBTITLE_PLACEHOLDERS:
            role = "subtitle"
        elif ph:
            role = "body"
        else:
            role = "textbox"
        return [
            TextBlock(
                **_geometry(shape),
                paragraphs=paragraphs,
                role=role,
                placeholder=ph or None,
            )
        ]

    return []


def _slide_title(blocks: list[ShapeBase]) -> str:
    """标题来源：标题占位符优先；否则取「短且加粗」的首个段落（常见于模板外的手写标题）。"""
    for block in blocks:
        if isinstance(block, TextBlock) and block.role == "title":
            text = " ".join(p.text.strip() for p in block.paragraphs if p.text.strip())
            if text:
                return text
    for block in blocks:
        if not isinstance(block, TextBlock):
            continue
        for paragraph in block.paragraphs:
            text = paragraph.text.strip().replace("\n", " ")
            if not text or len(text) > 24:
                continue
            if paragraph.runs and all(run.bold for run in paragraph.runs if run.text.strip()):
                return text
    return ""


def _notes(slide: Any) -> list[Paragraph]:
    try:
        if not slide.has_notes_slide:
            return []
        frame = slide.notes_slide.notes_text_frame
    except Exception:  # pragma: no cover
        return []
    paragraphs = _paragraphs(frame)
    return [p for p in paragraphs if not p.is_empty]


def read_pptx(path: str | Path, sink: AssetSink) -> Deck:
    """解析 PPTX，图片/媒体会即时写入 ``sink`` 指向的目录。"""
    source = Path(path)
    if source.suffix.lower() == ".ppt":
        raise ValueError(
            f"{source.name}: 不支持旧版二进制 .ppt，请先用 PowerPoint 或 LibreOffice 另存为 .pptx"
        )

    presentation = Presentation(str(source))
    deck = Deck(
        source=source,
        slide_width=int(presentation.slide_width or 0),
        slide_height=int(presentation.slide_height or 0),
    )

    for index, slide in enumerate(presentation.slides, start=1):
        blocks: list[ShapeBase] = []
        for shape in reading_order(slide.shapes):
            blocks.extend(_shape_blocks(shape, index, sink))
        deck.slides.append(
            Slide(number=index, title=_slide_title(blocks), blocks=blocks, notes=_notes(slide))
        )

    core_title = ""
    try:
        core_title = (presentation.core_properties.title or "").strip()
    except Exception:  # pragma: no cover
        core_title = ""

    deck.title = (
        core_title
        or (deck.slides[0].title if deck.slides and deck.slides[0].title else "")
        or source.stem
    )
    return deck
